from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = r"C:\OpenVLA-Simulator\4.4.4 虚实迁移实验结果（实际仿真数据）.docx"


def set_font(run, size=12, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(2.54)
section.right_margin = Cm(2.54)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(14)
run = title.add_run("4.4.4 虚实迁移实验结果")
set_font(run, size=16, bold=True)

add_body(doc, "为验证虚实迁移中可在仿真环境独立评估的模块效果，本文以本地 OpenVLA-7B 的 LIBERO-Goal 微调权重为策略基线，在 MuJoCo/LIBERO 中构造源域与目标视觉域。源域使用原始相机图像；目标域对每帧图像施加固定 RGB 通道增益（0.72、1.05、0.82）、偏置（0.04）和标准差为 0.02 的高斯观测噪声，以模拟相机曝光、白平衡和传感器噪声造成的外观差异。")
add_body(doc, "实验选取“将碗放至炉灶”“将酒瓶放至橱柜顶部”“将餐盘推至炉灶前方”和“打开炉灶”4 个 LIBERO-Goal 任务。各任务使用同一随机种子（20260728）和 1 个固定初始状态，回合最多执行 310 个环境步（前 10 步为夹爪初始化，最多 300 个 OpenVLA 决策步）；不使用专家动作、真实状态恢复、人工接管或实体机器人数据，成功与否完全由环境任务谓词判定。")

caption = doc.add_paragraph()
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_before = Pt(6)
caption.paragraph_format.space_after = Pt(4)
run = caption.add_run("表 4-4 虚实迁移仿真实验的实际运行结果")
set_font(run, size=11, bold=True)

table = doc.add_table(rows=1, cols=5)
table.style = "Table Grid"
headers = ["评测条件", "成功数/总数", "成功率", "失败任务", "说明"]
for cell, text in zip(table.rows[0].cells, headers):
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "D9EAF7")
    for run in cell.paragraphs[0].runs:
        set_font(run, size=10.5, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

rows = [
    ["源域（原始相机图像）", "4/4", "100.0%", "无", "OpenVLA 基线"],
    ["目标域（视觉扰动）", "3/4", "75.0%", "将餐盘推至炉灶前方", "未见相机外观变化"],
    ["目标域 + 光度校准模块", "4/4", "100.0%", "无", "反演通道增益与偏置；噪声保留"],
]
for values in rows:
    cells = table.add_row().cells
    for index, (cell, value) in enumerate(zip(cells, values)):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_font(run, size=10.5)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT

add_body(doc, "由表4-4可见，原始 OpenVLA 在源域中完成4个任务；引入目标域相机扰动后，推盘任务在300个决策步内未完成，成功率降至75.0%。加入视觉光度校准模块后，4个任务均成功完成，目标域成功率恢复至100.0%，较未校准的目标域提高25.0个百分点。该结果表明，在本实验构造的外观差异下，输入光度校准能够减少相机响应变化对视觉—语言—动作策略的影响。")
add_body(doc, "需要说明的是，本节结果来自4个固定初始状态的实际仿真回合，属于小样本受控验证；它证明的是所设置视觉扰动下的仿真迁移有效性，而非实体机器人上的 Sim-to-Real 成功率。原始逐回合记录保存在 outputs/4_4_4_actual_transfer/source_310.json、target_shifted_310.json 与 target_calibrated_310.json。")

doc.save(OUT)
print(OUT)
