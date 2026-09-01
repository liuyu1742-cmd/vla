from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\OpenVLA-Simulator")
SOURCE = ROOT / "outputs" / "midterm_testing_vla82" / "VLA82范围内_中期8类60物体测试执行台账_真实闭环更新.xlsx"
OUTPUT = ROOT / "outputs" / "midterm_testing_vla82" / "VLA82范围内_中期8类60物体技能与操作类型汇总.docx"


def set_run_font(run, name="微软雅黑", size=9, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="9E9E9E", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_grid(table, widths_cm):
    tbl = table._tbl
    grid = tbl.tblGrid
    grid_cols = grid.findall(qn("w:gridCol"))
    while len(grid_cols) < len(widths_cm):
        col = OxmlElement("w:gridCol")
        grid.append(col)
        grid_cols.append(col)
    for col, width_cm in zip(grid_cols, widths_cm):
        col.set(qn("w:w"), str(int(width_cm * 567)))

    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def repeat_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_no_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def configure_paragraph(paragraph, alignment=None, space_before=0, space_after=0, line_spacing=1.05):
    if alignment is not None:
        paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing


def write_cell(cell, text, alignment=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=9, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    configure_paragraph(paragraph, alignment=alignment, line_spacing=1.05)
    run = paragraph.add_run("" if text is None else str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def read_records():
    wb = openpyxl.load_workbook(SOURCE, data_only=True, read_only=True)
    task_rows = list(wb["8类任务"].iter_rows(values_only=True))
    object_rows = list(wb["60种物体"].iter_rows(values_only=True))
    task_headers = list(task_rows[0])
    object_headers = list(object_rows[0])
    task_index = {str(v): i for i, v in enumerate(task_headers) if v is not None}
    object_index = {str(v): i for i, v in enumerate(object_headers) if v is not None}

    category_order = []
    expected_counts = {}
    for row in task_rows[1:]:
        category = row[task_index["任务类别"]]
        if category:
            category_order.append(category)
            expected_counts[category] = int(row[task_index["入选物体数"]])

    records = []
    for row in object_rows[1:]:
        if not row[object_index["选择ID"]]:
            continue
        records.append(
            {
                "category": row[object_index["任务类别"]],
                "object": row[object_index["物体"]],
                "skill": row[object_index["表内操作标签"]],
                "operation": row[object_index["真实视频标注"]],
            }
        )

    grouped = {category: [] for category in category_order}
    for record in records:
        if record["category"] not in grouped:
            raise ValueError(f"发现未列入8类任务的类别：{record['category']}")
        grouped[record["category"]].append(record)

    if len(category_order) != 8:
        raise ValueError(f"任务类别数不是8：{len(category_order)}")
    if len(records) != 60:
        raise ValueError(f"物体记录数不是60：{len(records)}")
    for category in category_order:
        actual = len(grouped[category])
        if actual != expected_counts[category]:
            raise ValueError(f"{category}数量不一致：期望{expected_counts[category]}，实际{actual}")

    return category_order, grouped, expected_counts


def add_category_table(doc, category, records, index, total):
    heading = doc.add_paragraph()
    configure_paragraph(heading, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=2, space_after=5)
    heading.paragraph_format.keep_with_next = True
    heading_run = heading.add_run(f"{index}. {category}（{len(records)}项）")
    set_run_font(heading_run, size=12, bold=True, color=(31, 78, 121))

    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    table.allow_autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    widths = [1.5, 10.5, 3.8, 11.0]
    set_table_grid(table, widths)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    headers = ["序号", "技能名称", "控制对象", "操作类型"]
    header_row = table.rows[0]
    repeat_header_row(header_row)
    set_no_row_split(header_row)
    for cell, width, header in zip(header_row.cells, widths, headers):
        set_cell_width(cell, width)
        set_cell_shading(cell, "D9E2F3")
        write_cell(cell, header, alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=9, color=(31, 31, 31))

    for serial, record in enumerate(records, 1):
        row = table.add_row()
        set_no_row_split(row)
        values = [serial, record["skill"], record["object"], record["operation"]]
        alignments = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]
        for cell, width, value, alignment in zip(row.cells, widths, values, alignments):
            set_cell_width(cell, width)
            write_cell(cell, value, alignment=alignment, size=9)

    if index < total:
        doc.add_page_break()


def build_doc():
    category_order, grouped, _ = read_records()

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(9)

    title = doc.add_paragraph()
    configure_paragraph(title, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    title_run = title.add_run("VLA82范围内中期8类60物体技能与操作类型汇总")
    set_run_font(title_run, size=16, bold=True, color=(31, 78, 121))

    note = doc.add_paragraph()
    configure_paragraph(note, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    note_run = note.add_run("数据来源：VLA82范围内_中期8类60物体测试执行台账_真实闭环更新.xlsx；技能名称取“表内操作标签”，操作类型取“真实视频标注”。")
    set_run_font(note_run, size=8.5, color=(89, 89, 89))

    for index, category in enumerate(category_order, 1):
        add_category_table(doc, category, grouped[category], index, len(category_order))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"saved={OUTPUT}")
    print(f"categories={len(category_order)} records={sum(len(v) for v in grouped.values())}")


if __name__ == "__main__":
    build_doc()
