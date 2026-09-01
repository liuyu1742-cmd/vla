from copy import deepcopy
from pathlib import Path
from shutil import copy2
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(r"C:\OpenVLA-Simulator")
SRC=ROOT/"outputs/evidence_revised_report/evidence_report_chapter5_academic_v3.docx"
REF=Path(r"C:\Users\sjtu101\Desktop/2. 基于视觉的人类动作捕获、解析和技能学习迁移技术研究.docx")
OUT=ROOT/"outputs/evidence_revised_report/evidence_report_chapter5_academic_v4.docx"
FIG=ROOT/"outputs/formal_skill_augmented_r2_eval/probes/seed_101_frame0.png"
def find(d,p):
 m=[x for x in d.paragraphs if x.text.strip().startswith(p)]
 return m[-1]
def after(d,p):
 ps=d.paragraphs
 for i,x in enumerate(ps):
  if x._p is p._p:return ps[i+1]
def fmt(s,t):
 if s._p.pPr is not None:
  q=t._p.get_or_add_pPr();q.clear();q.append(deepcopy(s._p.pPr))
 if s.runs and t.runs and s.runs[0]._element.rPr is not None:
  q=t.runs[0]._element.get_or_add_rPr();q.clear();q.append(deepcopy(s.runs[0]._element.rPr))
def add(p,text,style):
 x=OxmlElement('w:p');p._p.addnext(x);q=p._parent.add_paragraph();q._p.getparent().remove(q._p);q._p=x;q._element=x;q.add_run(text);fmt(style,q);return q
def remove_between(a,b):
 n=a._p.getnext()
 while n is not b._p:
  nxt=n.getnext();n.getparent().remove(n);n=nxt
def main():
 copy2(SRC,OUT);d=Document(OUT);r=Document(REF);h=find(d,'5.4 ');h5=find(d,'5.5 ');style=after(r,find(r,'5.2.1 '));cap=find(r,'图 5.1 ');fmt(find(r,'5.4 '),h);remove_between(h,h5)
 p=add(h,'技能检索与动态组合是技能库由目录化知识走向任务执行的关键环节。对于用户提出的室内服务请求，系统并不直接将自然语言映射为单一动作，而是先识别其中包含的任务意图、目标对象、目标区域和限制条件，再从第5.2节的任务分类与第5.3节的对象操作目录中选择相应条目。例如，“将水杯放到托盘上”同时涉及水杯定位、稳定抓取、室内搬运和托盘区域放置；“整理书桌”则需要围绕书籍、文具和收纳容器检索多个对象操作。由此，检索结果必须同时满足任务语义、对象标签和场景约束，而不能仅依赖关键词匹配。',style)
 p=add(p,'为统一描述查询与技能条目，设用户请求为q=(t_q,O_q,G_q,C_q)，其中t_q表示任务意图，O_q表示目标对象集合，G_q表示目标区域，C_q表示约束条件；设候选技能为s_i=(t_i,O_i,A_i,G_i,C_i,E_i)，其中A_i为动作原语序列，E_i为该条目的示教证据与状态信息。检索评分需要综合衡量任务类别相似度、对象与区域匹配程度以及约束条件的一致性。令sim_t表示任务语义相似度，J(·)表示集合的Jaccard匹配系数，sim_c表示场景约束相似度，λ_t、λ_o、λ_g、λ_c为满足和为1的权重，则候选技能的综合评分定义为：',style)
 p=add(p,'Score(q,s_i)=λ_t·sim_t(t_q,t_i)+λ_o·J(O_q,O_i)+λ_g·J(G_q,G_i)+λ_c·sim_c(C_q,C_i)，  Σλ=1。',style);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p=add(p,'按照上述评分，系统先获得候选技能的排序，再结合当前场景观测筛除对象不可见、目标区域不可达、容器状态不满足或安全条件不成立的条目。对象、场景和约束的联合筛选使技能检索能够保持与项目对象操作清单的一致性，并为后续的动作组合提供明确的输入。',style)
 p=add(p,'当一项服务请求涉及多个对象或多个状态变化时，系统需要将检索到的原子技能组织为有向执行序列。设组合计划为Π=(s_{i1},s_{i2},…,s_{im})，x_k为第k步执行前的环境状态，Pre(s_{ik})为第k个技能的前置条件，T(x_k,s_{ik})为该技能引起的状态转移。只有当目标对象已定位、容器处于可操作状态、目标区域未被占用且路径满足安全边界时，相应技能才能进入序列；执行后生成的新状态又成为后续技能的输入。因此，组合过程不仅追求候选技能的相关性，还需要避免碰撞、跌落、液体倾倒、热源接近和设备误触等风险。以Risk(s_{ik},x_k)表示第k步在当前状态下的风险代价，μ表示风险惩罚系数，则最优组合计划写为：',style)
 p=add(p,'Π*=arg max_Π Σ_k Score(q,s_{ik})－μ·Σ_k Risk(s_{ik},x_k)，  s.t. Pre(s_{ik})⊆x_k，x_{k+1}=T(x_k,s_{ik})。',style);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p=add(p,'该优化表达说明：组合计划既应由高相关度技能构成，也应满足对象状态和时序约束。对于收纳、递送和清洁等多步骤任务，系统可先安排无遮挡、可达且风险较低的对象操作，再执行依赖容器开合、区域清空或前一步放置结果的后续操作。图5.3给出了项目RoboCasa取放场景中的对象与目标区域示例，用于说明技能条目在具体仿真环境中的语义落点。',style)
 q=add(p,'',style);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run().add_picture(str(FIG),width=Cm(13.4));c=add(q,'图 5.3 RoboCasa室内取放场景示例（项目输出）',cap);c.alignment=WD_ALIGN_PARAGRAPH.CENTER
 u=OxmlElement('w:updateFields');u.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val','true');d.settings.element.append(u);d.save(OUT);print(OUT)
if __name__=='__main__':main()
