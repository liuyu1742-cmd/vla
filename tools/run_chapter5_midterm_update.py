import copy
import importlib.machinery
import importlib.util
import subprocess
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

PYC = Path(r"tools/__pycache__/update_chapter5_midterm_simplified.cpython-312.pyc")
loader = importlib.machinery.SourcelessFileLoader("midterm", str(PYC))
spec = importlib.util.spec_from_loader("midterm", loader)
m = importlib.util.module_from_spec(spec)
loader.exec_module(m)


def extract_figures(dataset_root: Path, figure_dir: Path):
    figure_dir.mkdir(parents=True, exist_ok=True)
    figures = {}
    for code, (source_code, obj) in m.FIGURE_SOURCES.items():
        obj_dir = m._task_dir(dataset_root, source_code) / obj
        video = next(obj_dir.rglob("*.mp4"))
        jpg = figure_dir / f"{code.replace('.', '_')}_{obj}.jpg"
        png = jpg.with_suffix(".png")
        subprocess.run(["ffmpeg", "-y", "-ss", "0", "-i", str(video), "-frames:v", "1", "-q:v", "2", str(jpg)], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        with Image.open(jpg) as image:
            image.convert("RGB").save(png, format="PNG")
        figures[code] = png
    return figures


def replace_sections(doc, figures):
    network_codes = {"5.2.1", "5.2.2", "5.2.3", "5.2.4", "5.2.7"}
    for ordinal, code in enumerate(m.TASK_TITLES, start=1):
        heading = m._heading(doc, code)
        heading.clear()
        heading.add_run(f"{code} {m.TASK_TITLES[code]}")
        body = m._clear_between(doc, heading)
        body.clear()
        body.add_run(m.SECTION_TEXT[code])
        figure_p = m._insert_after(body)
        figure_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure_p.add_run().add_picture(str(figures[code]), width=Cm(11.5))
        _, obj = m.FIGURE_SOURCES[code]
        if code in network_codes:
            caption_text = f"\u56fe5-{ordinal} {m.TASK_TITLES[code]}\u76f8\u5173\u8bbe\u5907\u89c6\u89c9\u4e0a\u4e0b\u6587\u5e27\uff08\u5bf9\u8c61\uff1a{obj}\uff1b\u7f51\u7edc\u63a7\u5236\u4ee5\u63a5\u53e3\u65e5\u5fd7\u9a8c\u6536\uff09"
        else:
            caption_text = f"\u56fe5-{ordinal} {m.TASK_TITLES[code]}\u76f8\u5173VLA\u793a\u6559\u5e27\uff08\u5bf9\u8c61\uff1a{obj}\uff09"
        caption = m._insert_after(figure_p, caption_text)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER


def caption_before(doc, element, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    element.addprevious(p._p)


def insert_original_tables(doc, original):
    reference = doc.tables[1]
    rows = [
        ("\u88685-2 \u5bb6\u7535\u63a7\u5236\u7c7b\u6280\u80fd\uff0840\u79cd\uff09", 18),
        ("\u88685-3 \u73af\u5883\u8c03\u8282\u7c7b\u6280\u80fd\uff0820\u79cd\uff09", 19),
        ("\u88685-4 \u5b89\u9632\u76d1\u63a7\u7c7b\u6280\u80fd\uff0815\u79cd\uff09", 20),
        ("\u88685-7 \u517b\u62a4\u7ba1\u7406\u7c7b\u6280\u80fd\uff0810\u79cd\uff09", 23),
    ]
    for caption, index in rows:
        table = copy.deepcopy(original.tables[index]._tbl)
        reference._tbl.addprevious(table)
        caption_before(doc, table, caption)
    caption_before(doc, reference._tbl, "\u88685-8 \u5df2\u6838\u9a8cVLA\u7269\u4f53\u64cd\u4f5c\u6837\u672c\uff0888\u79cd\uff09")


def update(doc_path, original_path, dataset_root, figure_dir):
    doc = Document(doc_path)
    original = Document(original_path)
    figures = extract_figures(dataset_root, figure_dir)
    m._rewrite_summary(doc)
    replace_sections(doc, figures)
    insert_original_tables(doc, original)
    doc.save(doc_path)


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--doc", type=Path, required=True)
    parser.add_argument("--original", type=Path, required=True)
    parser.add_argument("--dataset-root", type=Path, required=True)
    parser.add_argument("--figure-dir", type=Path, required=True)
    args = parser.parse_args()
    update(args.doc, args.original, args.dataset_root, args.figure_dir)
    print("updated")
