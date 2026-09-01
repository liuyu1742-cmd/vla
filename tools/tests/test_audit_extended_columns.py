from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from tools.audit_chapter5_object_alignment import load_document_objects


def test_load_document_objects_accepts_extended_five_column_audit_table():
    with TemporaryDirectory() as temp_dir:
        path = Path(temp_dir) / "audit.docx"
        doc = Document()
        doc.add_table(rows=1, cols=1)
        table = doc.add_table(rows=2, cols=5)
        for cell, text in zip(table.rows[0].cells, ["技能名称", "操作对象", "操作说明", "源指令", "核验状态"]):
            cell.text = text
        for cell, text in zip(table.rows[1].cells, ["物品递送", "书籍", "床上至床头柜", "Move book", "完整"]):
            cell.text = text
        doc.save(path)
        objects, rows = load_document_objects(path)

    assert objects["5.2.8"] == {"书籍"}
    assert rows[0]["document_operation"] == "床上至床头柜"
