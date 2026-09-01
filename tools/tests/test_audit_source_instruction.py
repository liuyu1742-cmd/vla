from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from tools.audit_chapter5_object_alignment import load_document_objects


def test_load_document_objects_retains_exact_source_instruction_column():
    with TemporaryDirectory() as temp_dir:
        path = Path(temp_dir) / "audit.docx"
        doc = Document()
        doc.add_table(rows=1, cols=1)
        table = doc.add_table(rows=2, cols=5)
        for cell, text in zip(table.rows[1].cells, ["物品递送", "书籍", "床上至床头柜", "Move book from bed to nightstand", "完整"]):
            cell.text = text
        doc.save(path)
        _, rows = load_document_objects(path)

    assert rows[0]["document_source_instruction"] == "Move book from bed to nightstand"
