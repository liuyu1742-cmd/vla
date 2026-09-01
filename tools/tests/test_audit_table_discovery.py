import importlib.util
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

MODULE = Path(__file__).parents[1] / "audit_chapter5_object_alignment.py"
SPEC = importlib.util.spec_from_file_location("audit_extra", MODULE)
MOD = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MOD)


def test_load_document_objects_finds_audit_table_after_reference_tables():
    with TemporaryDirectory() as temp_dir:
        path = Path(temp_dir) / "doc.docx"
        doc = Document()
        doc.add_table(rows=1, cols=2)
        doc.add_table(rows=1, cols=2)
        audit = doc.add_table(rows=2, cols=5)
        headers = ["技能名称", "操作对象", "操作说明", "源指令", "核验状态"]
        for cell, text in zip(audit.rows[0].cells, headers):
            cell.text = text
        values = [next(iter(MOD.SKILL_TO_TASK)), "物体", "放置", "put object", "视频+动作"]
        for cell, text in zip(audit.rows[1].cells, values):
            cell.text = text
        doc.save(path)
        grouped, rows = MOD.load_document_objects(path)
        assert grouped == {next(iter(MOD.SKILL_TO_TASK.values())): {"物体"}}
        assert rows[0]["document_source_instruction"] == "put object"


