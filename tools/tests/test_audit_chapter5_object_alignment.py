import unittest


class TestObjectAlignment(unittest.TestCase):
    def test_compare_object_sets_reports_document_only_and_dataset_only_objects(self):
        from tools.audit_chapter5_object_alignment import compare_object_sets

        result = compare_object_sets(
            {"整理收纳": {"收纳盒", "收纳箱", "钢笔"}},
            {"整理收纳": {"收纳盒", "收纳箱", "钢笔", "笔盒"}},
        )

        self.assertEqual(result["document_only"], {"整理收纳": ["笔盒"]})
        self.assertEqual(result["dataset_only"], {})


if __name__ == "__main__":
    unittest.main()

    def test_load_document_objects_accepts_extended_five_column_audit_table(self):
        from pathlib import Path
        from tempfile import TemporaryDirectory
        from docx import Document
        from tools.audit_chapter5_object_alignment import load_document_objects

        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "audit.docx"
            doc = Document()
            doc.add_table(rows=1, cols=1)
            table = doc.add_table(rows=2, cols=5)
            for cell, text in zip(table.rows[0].cells, ["技能名称", "操作对象", "操作说明", "源指令", "核验状态"]):
                cell.text = text
            for cell, text in zip(table.rows[1].cells, ["物品递送", "书籍", "床上至床头柜", "Move book", "完整"]):
                cell.text = text
            doc.save(path)
            objects, rows = load_document_objects(path)

        self.assertEqual(objects["5.2.8"], {"书籍"})
        self.assertEqual(rows[0]["document_operation"], "床上至床头柜")