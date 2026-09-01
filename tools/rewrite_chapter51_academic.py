from copy import deepcopy
from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

ROOT=Path(r"C:\OpenVLA-Simulator")
SRC=ROOT/"outputs/evidence_revised_report/evidence_report_chapter5_academic.docx"
REF=Path(r"C:\Users\sjtu101\Desktop/2. 基于视觉的人类动作捕获、解析和技能学习迁移技术研究.docx")
OUT=ROOT/"outputs/evidence_revised_report/evidence_report_chapter5_academic_v2.docx"
FIG=ROOT/"outputs/evidence_revised_report/figure_5_1_academic_framework.png"
def find(doc,prefix):
 m=[p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
 if not m: raise RuntimeError(prefix)
 return m[-1]
def after(doc,p):
 ps=doc.paragraphs
 for i,x in enumerate(ps):
  if x._p is p._p:return ps[i+1]
def fmt(src,dst):
 if src._p.pPr is not None:
  q=dst._p.get_or_add_pPr();q.clear();q.append(deepcopy(src._p.pPr))
 if src.runs and dst.runs and src.runs[0]._element.rPr is not None:
  q=dst.runs[0]._element.get_or_add_rPr();q.clear();q.append(deepcopy(src.runs[0]._element.rPr))
def insert_after(p,text,style):
 x=OxmlElement('w:p');p._p.addnext(x);q=p._parent.add_paragraph();q._p.getparent().remove(q._p);q._p=x;q._element=x
 if text:q.add_run(text)
 fmt(style,q);return q
def remove_between(first,last):
 n=first._p.getnext()
 while n is not last._p:
  nxt=n.getnext();n.getparent().remove(n);n=nxt
def main():
 copy2(SRC,OUT);d=Document(OUT);r=Document(REF)
 rh1=find(r,'5.1 ');rh3=find(r,'5.2.1 ');body=after(r,rh3);cap=find(r,'图 5.1 ')
 h1=find(d,'5.1 ');h2=find(d,'5.2 ');remove_between(h1,h2);fmt(rh1,h1)
 p=insert_after(h1,'技能库是面向家庭服务机器人多任务执行的技能知识管理系统，其基本目标是在统一框架下完成技能条目的规范存储、按需检索、约束校验与组合调用。与仅保存动作轨迹或单条指令样本的资源库不同，本研究中的技能库以“任务目标—对象交互—状态变化”为核心组织单位：任务目标说明服务意图和适用场景，对象交互刻画机器人与物体、容器、设备或区域之间的关系，状态变化给出操作前后可观测的完成条件。由此，技能库既可记录“书籍装箱并归位”“水杯放置到托盘”等具体能力，也能够表达开合、抓取、放置、旋转、按压和擦拭等可在不同任务中复用的基础操作。该设计使前四章形成的视觉感知、动作理解和策略执行能力能够以结构化技能单元的形式进入后续家庭服务任务。',body)
 p=insert_after(p,'技能库的数据模型采用“任务技能—对象操作—动作原语—参数与约束—证据索引”的分层表示。任务技能层定义服务类别、场景边界、目标状态和技能名称；对象操作层记录控制对象、目标区域及对象状态变化，例如将水杯由桌面移动至托盘、将收纳箱由打开状态转变为关闭状态；动作原语层描述接近、抓取、搬运、放置、开合、旋转、按压、推拉和擦拭等基本操作；参数与约束层记录动作所需的目标位姿、容器朝向、接触方式、可达空间、速度或安全边界；证据索引层关联公开示教来源、本地对象操作清单和仿真场景记录。该模型兼顾任务层的语义表达与执行层的可操作性，支持将复杂家庭服务拆分为前后依赖明确的原子操作，也支持在不同对象和场景之间复用相同动作原语。',body)
 p=insert_after(p,'技能库由技能元数据、对象操作目录、约束与状态描述、证据索引以及检索组合接口五部分组成。技能元数据记录任务类别、适用区域、输入输出和版本信息；对象操作目录覆盖本项目15类室内任务及120种对象物体，采用“任务类别：对象+操作”的统一命名；约束与状态描述用于保存容器开合、物体姿态、目标区域占用和安全限制等条件；证据索引将条目与HABIT、BEHAVIOR-1K、DROID等公开真实示教的核验状态及RoboCasa场景联系起来；检索组合接口则根据用户指令、对象标签和场景条件返回候选技能及其执行顺序。在管理功能上，技能条目支持新增、修订、查询、版本保留和证据状态更新，从而保证技能扩展过程可追溯、可复核，并为第5.4节的检索与动态组合提供统一的数据基础。',body)
 q=insert_after(p,'',body);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run().add_picture(str(FIG),width=Cm(14.1))
 c=insert_after(q,'图 5.1 项目技能库总体架构',cap);c.alignment=WD_ALIGN_PARAGRAPH.CENTER
 u=OxmlElement('w:updateFields');u.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val','true');d.settings.element.append(u)
 d.save(OUT);print(OUT)
if __name__=='__main__':main()
