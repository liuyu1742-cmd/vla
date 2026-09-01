"""Build the expert-facing VLA82 midterm simulation test report."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "midterm_testing_vla82" / "reports" / "VLA82中期仿真测试报告_专家版.docx"

NAVY = "1F4D78"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
INK = RGBColor(31, 37, 45)


def set_font(run, *, name: str = "Microsoft YaHei", size: float = 10.5, bold: bool = False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color or INK


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for index, width in enumerate(widths):
        if index < len(grid.gridCol_lst):
            grid.gridCol_lst[index].set(qn("w:w"), str(width))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def para(doc, text="", *, style=None, size=10.5, bold=False, color=None, align=None, before=0, after=6, line=1.1):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    set_font(p.add_run(text), size=10.5)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 11.5}[level], bold=True, color=RGBColor.from_string(BLUE if level < 3 else NAVY))
    return p


def add_image(doc, image: Path, caption: str, width=5.75):
    if not image.is_file():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(image), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    set_font(cap.add_run(caption), size=9, color=RGBColor(89, 89, 89))


def add_two_images(doc, left: Path, right: Path, left_caption: str, right_caption: str):
    table = doc.add_table(rows=2, cols=2)
    set_table_widths(table, [4680, 4680])
    for cell in table.rows[0].cells:
        cell.text = ""
    for image, cell in zip((left, right), table.rows[0].cells):
        if image.is_file():
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(image), width=Inches(3.02))
    for text, cell in zip((left_caption, right_caption), table.rows[1].cells):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), size=8.5, color=RGBColor(89, 89, 89))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_result_table(doc):
    rows = [
        ("VLA82-001", "垃圾桶投放", "抓取—跨区域转运—容器内稳定释放", "PASS / predicate_success=true", "402"),
        ("VLA82-002", "典型物体操作", "抓取、转运与目标区域稳定释放", "PASS / predicate_success=true", "已保存过程记录"),
        ("VLA82-004、005", "抓取/放置", "物体接触、抬升、转运与释放", "PASS / predicate_success=true", "已保存过程记录"),
        ("VLA82-007", "设备控制", "与指定旋钮接触后完成关节旋转", "PASS / predicate_success=true", "已保存过程记录"),
        ("VLA82-017、019、020", "柜体相关物体操作", "操作序列与稳定放置闭环", "PASS / predicate_success=true", "600+"),
        ("VLA82-030", "设备/抽屉关闭", "指定夹爪接触、滑动关节变化与闭合终态", "PASS / predicate_success=true", "74"),
        ("VLA82-036、038、042", "代表性家政对象", "源任务阶段与物理终态联合判定", "PASS / predicate_success=true", "已保存过程记录"),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    set_table_widths(table, [1250, 1550, 3000, 2200, 1360])
    headers = ("编号", "操作类别", "严格判定要点", "判定结果", "轨迹步数")
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), size=9, bold=True, color=RGBColor.from_string(NAVY))
    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (0, 3, 4) else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(text), size=8.7)
    return table


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("VLA82 中期仿真测试报告"), size=8.5, color=RGBColor(89, 89, 89))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("OpenVLA-Simulator  •  测试证据基于可追溯轨迹与物理判据"), size=8, color=RGBColor(110, 110, 110))

    para(doc, "中期仿真测试报告", size=23, bold=True, color=RGBColor.from_string(NAVY), before=10, after=3)
    para(doc, "视觉模仿学习与机器人家政操作复现", size=13, color=RGBColor(70, 70, 70), after=14)
    meta = doc.add_table(rows=3, cols=2)
    set_table_widths(meta, [1800, 7560])
    for row, (label, value) in zip(meta.rows, (("测试日期", "2026年8月"), ("测试平台", "OpenVLA-Simulator / RoboCasa / MuJoCo / PandaOmron"), ("报告定位", "中期阶段：测试步骤、代表性结果与过程图片"))):
        shade(row.cells[0], LIGHT_GRAY)
        set_font(row.cells[0].paragraphs[0].add_run(label), size=9.5, bold=True, color=RGBColor.from_string(NAVY))
        set_font(row.cells[1].paragraphs[0].add_run(value), size=9.5)

    para(doc, "执行摘要", size=12.5, bold=True, color=RGBColor.from_string(BLUE), before=15, after=5)
    para(doc, "本报告汇总视觉模仿学习、技能序列执行和机器人仿真复现的中期测试工作。测试按 8 类典型家政任务组织，并以 VLA82 范围内 60 种典型物体建立逐物体测试记录。测试采用真实视频操作标注驱动的操作规格，所有控制命令均经由仿真环境公开 step 接口执行；成功结论由接触、物体状态、设备关节变化和稳定终态等物理观测共同给出。当前已形成 12 项代表性对象的严格 PASS 闭环轨迹，覆盖抓取/放置、容器投放、柜体相关操作及设备控制等关键操作范式。")

    heading(doc, "1. 测试目标与判定依据", 1)
    para(doc, "中期测试围绕“视觉设备获取示范—技能表示生成—机器人仿真复现—物理结果判定”的闭环组织。项目目标范围为 8 类典型家政任务与 60 种典型物体；每一物体均关联唯一编号、任务类别、真实视频操作标注、仿真资产映射、过程图像和判定记录。本报告重点展示当前已完成的正向验证结果。")
    add_bullet(doc, "技能学习：由源任务文本、视觉观察和动作阶段生成可执行的技能序列（Skill IR/阶段控制）。")
    add_bullet(doc, "机器人复现：PandaOmron 在 RoboCasa/MuJoCo 场景中执行抓取、清洁相关接触、容器投放与设备控制动作。")
    add_bullet(doc, "严格判定：不以任务标签或控制器结束信号作为成功依据；要求观测到指定对象/夹爪接触、相应状态变化和稳定终态。")

    heading(doc, "2. 测试环境与数据记录", 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_widths(table, [2200, 7160])
    for cell, text in zip(table.rows[0].cells, ("要素", "配置与用途")):
        shade(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), size=9.5, bold=True, color=RGBColor.from_string(NAVY))
    for left, right in (
        ("仿真平台", "RoboCasa/MuJoCo 物理仿真；PandaOmron 移动操作机器人；完整厨房/家政场景。"),
        ("视觉输入", "主相机与腕部相机记录操作过程；重生成证据视频采用无遮挡宽视角以同时呈现机器人、对象与操作空间。"),
        ("动作执行", "技能控制器输出公开仿真动作，按源任务的 grasp、place、wipe、push、close、turn_knob 等阶段执行。"),
        ("物理核验", "保存每步物理快照、对象位姿、夹爪/对象接触、设备关节位置及稳定释放等观测量。"),
        ("可追溯产物", "每条严格 PASS 轨迹均保存 episode JSON（判定与诊断）和 NPZ（图像、动作、快照、谓词结果）。"),
    ):
        cells = table.add_row().cells
        set_font(cells[0].paragraphs[0].add_run(left), size=9, bold=True, color=RGBColor.from_string(NAVY))
        set_font(cells[1].paragraphs[0].add_run(right), size=9)

    heading(doc, "3. 测试流程", 1)
    steps = (
        ("01", "加载任务与对象", "从 VLA82 数据范围解析源任务、典型物体和真实视频操作标注，生成操作规格。"),
        ("02", "构建仿真场景", "实例化指定物体/设备部件、目标区域、相机与 PandaOmron 机器人。"),
        ("03", "技能阶段执行", "按操作阶段调度抓取、转运、释放、设备推拉或旋钮控制等动作。"),
        ("04", "在线物理采集", "逐步记录图像、动作、对象与夹爪接触、关节状态、物体位姿及目标支撑状态。"),
        ("05", "严格谓词判定", "以接触先后、状态变化、终态稳定性和目标关系构成联合成功判定。"),
        ("06", "证据固化", "输出 JSON、NPZ、MP4 与截图；视频与源严格 PASS 轨迹通过侧车 JSON 建立关联。"),
    )
    flow = doc.add_table(rows=1, cols=3)
    flow.style = "Table Grid"
    set_table_widths(flow, [720, 2400, 6240])
    for cell, text in zip(flow.rows[0].cells, ("步骤", "阶段", "执行内容")):
        shade(cell, LIGHT_BLUE); cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(cell.paragraphs[0].add_run(text), size=9, bold=True, color=RGBColor.from_string(NAVY))
    for code, name, detail in steps:
        cells = flow.add_row().cells
        for index, value in enumerate((code, name, detail)):
            cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_font(cells[index].paragraphs[0].add_run(value), size=8.8, bold=index == 0, color=RGBColor.from_string(NAVY) if index == 0 else INK)

    heading(doc, "3.0 8 类任务与 60 种物体的测试组织", 2)
    para(doc, "60 种物体按真实视频操作标注分配到 8 类家政任务中，测试时不改变标注中的主操作对象和目标关系。测试工作完成了 60 种物体的对象编号整理、真实视频操作文本关联、仿真资产映射、场景加载检查和逐物体过程记录；在此基础上，对代表性对象开展了严格物理闭环运行。8 类任务的具体测试内容如下。")
    para(doc, "（1）收纳投放：针对垃圾桶、收纳容器及待投放物体，执行抓取、抬升、跨区域转运和容器内释放；记录物体是否进入容器有效空间及释放后是否稳定。VLA82-001 完成了该类严格 PASS 运行。", before=2, after=4)
    para(doc, "（2）物品整理：针对台面、容器和日常小物体，执行定向抓取、搬运、摆放或归位；记录夹爪接触、抬升、转运距离和目标区域支撑接触。VLA82-002、004、005 等形成了代表性严格 PASS 记录。", before=0, after=4)
    para(doc, "（3）台面/表面清洁：针对清洁工具和待清洁表面，配置工具抓取、接触建立、沿表面扫动和覆盖范围记录；重点保存工具—表面接触持续时间、轨迹覆盖单元和清洁状态变化。", before=0, after=4)
    para(doc, "（4）餐具处理：针对餐具、清洗区域和放置区域，配置取放、插入/摆放及稳定释放动作；记录餐具与夹爪、目标支撑面或容器区域的接触关系。", before=0, after=4)
    para(doc, "（5）烹饪与加热辅助：针对烤箱、烤盘、锅具及相关物体，配置拿取、推入/取出、摆放和加热设备交互动作；记录设备部件的空间位置及与机器人末端的接触。", before=0, after=4)
    para(doc, "（6）食品存取：针对食品、冰箱/储物区域和台面，配置食品抓取、从储物空间取出或放入、台面放置等动作；记录物体在储物区域与台面之间的状态转换。", before=0, after=4)
    para(doc, "（7）柜体/抽屉操作：针对柜门、抽屉、把手和内部物体，配置接近把手、接触、开合或滑动以及物体取放动作；记录指定夹爪接触、滑动/开合关节变化和闭合终态。VLA82-017、019、020、030 形成了该类代表性过程视频和严格 PASS 记录。", before=0, after=4)
    para(doc, "（8）家电与设备控制：针对旋钮、按钮和可滑动部件，配置目标部件识别、末端对齐、接触后旋转/按压/推拉和终态保持；记录部件关节位置或状态量变化。VLA82-007 完成了旋钮控制的严格 PASS 运行。", before=0, after=6)
    para(doc, "逐物体测试时，系统先核验该物体或设备部件的真实 MuJoCo 几何、可见性与碰撞体，再加载对应任务场景。运行过程保存动作轨迹、双相机过程帧、物理快照和 JSON 判定记录；形成严格 PASS 的轨迹进一步导出 NPZ 和 MP4。")

    heading(doc, "3.1 任务规格解析与测试用例装载", 2)
    para(doc, "首先从 VLA82 对象范围、任务分类和真实视频操作标注中读取测试对象。系统将每一条标注解析为结构化操作规格，规格中包含任务编号、物体/设备部件、操作文本、操纵对象、目标关系和阶段序列。例如，投放任务被分解为“抓取—抬升—跨区域转运—容器内释放”，设备控制任务被分解为“接近—接触—关节驱动—终态保持”。该步骤的输出为与源任务一一对应的操作规格文件，作为后续场景生成、控制器调度和判定器选择的唯一输入。")

    heading(doc, "3.2 场景初始化与视觉观测配置", 2)
    para(doc, "依据操作规格在 RoboCasa/MuJoCo 中加载 PandaOmron、目标物体或设备部件、目标容器/支撑面以及必要的辅助物体。系统在重置后读取对象的实际 MuJoCo 几何、碰撞体、关节名称和目标区域边界，确认被操纵对象与源任务对象一致。视觉侧同步启用主相机和腕部相机；对于柜体、抽屉和垃圾桶等容易被局部结构遮挡的场景，采用宽视角对准“机器人末端—对象—目标区域”共同操作空间，以便在视频中直接复核机器人和目标部件的位置关系。")

    heading(doc, "3.3 视觉模仿技能阶段执行", 2)
    para(doc, "控制器不以单一终点动作替代操作过程，而是按操作规格逐阶段执行。抓取类阶段先将末端移动至物体可达的预接近位姿，再在接触成立后闭合夹爪并验证物体随末端抬升；转运阶段保持已建立的抓取关系，将物体移动到源标注对应的目标区域；释放阶段要求夹爪打开后物体继续与目标支撑面或容器内壁保持稳定接触。对于推拉、关闭和旋钮控制等设备操作，控制器先对齐指定可操作部件，只有在夹爪与该部件的精确接触已经观测到后，才沿滑轨方向或旋转方向施加动作。")

    heading(doc, "3.4 在线物理采集与过程记录", 2)
    para(doc, "在每一个仿真步，测试程序同步保存主相机帧、腕部相机帧、机器人本体状态、动作向量、对象位姿、夹爪开度、接触对、设备关节位置和目标区域状态。这样既能够复原机器人执行过程，也能够验证“接触发生在状态变化之前”的因果顺序。例如，设备关闭操作需要先出现夹爪与指定把手/滑动部件的接触，再出现对应关节向闭合方向的连续变化；容器投放操作需要在物体进入有效容积后观测到夹爪释放和后续稳定。")

    heading(doc, "3.5 严格物理判定与证据固化", 2)
    para(doc, "轨迹执行完成后，判定器仅使用采集到的物理快照进行离线判定。抓取/放置任务检查夹爪—物体接触、抬升距离、转运距离、目标支撑接触和稳定释放；容器投放任务额外检查物体是否处于容器有效腔体内；设备控制任务检查指定接触体、真实关节变化幅度及终态阈值。只有全部阶段满足时，轨迹 JSON 中才会记录 status=PASS 且 predicate_success=true。随后将动作、双相机帧、快照及判定结果打包为 NPZ 和 JSON，并导出 MP4 与截图。")

    heading(doc, "4. 已完成的严格物理闭环结果", 1)
    para(doc, "表 1 汇总当前代表性严格 PASS 记录。每一条记录均满足 status=PASS 且 predicate_success=true；轨迹数据保存了动作、相机帧、物理快照与判定结果。")
    add_result_table(doc)
    para(doc, "结果解释：", size=10.5, bold=True, color=RGBColor.from_string(NAVY), before=8, after=3)
    add_bullet(doc, "抓取/放置类：判据要求夹爪—目标物体接触、有效抬升与转运，并在指定区域形成稳定释放或支撑接触。")
    add_bullet(doc, "容器投放类：除抓取和转运外，还要求对象位于容器有效腔体范围内并保持稳定。")
    add_bullet(doc, "设备控制类：要求夹爪与指定可操作部件产生接触，且真实滑动/旋转关节发生规定方向的变化并达到终态。")

    heading(doc, "5. 过程图像证据", 1)
    para(doc, "以下截图来自实际仿真过程或与严格 PASS 动作文件关联的同物理环境宽视角回放。取景策略以“机器人、对象/容器或设备部件、关键操作空间同时可见”为原则。")
    base = ROOT / "outputs" / "midterm_testing_vla82"
    add_two_images(
        doc,
        base / "video_regeneration_20260804_wide" / "VLA82-001" / "check_start.png",
        base / "video_regeneration_20260804_wide" / "VLA82-001" / "check_end.png",
        "图 1  VLA82-001：垃圾桶投放任务初始视角（容器与机器人完整入镜）",
        "图 2  VLA82-001：投放任务过程末段视角（容器位置持续可见）",
    )
    add_two_images(
        doc,
        base / "video_regeneration_20260804_clear" / "VLA82-017" / "replay_check_mid.png",
        base / "video_regeneration_20260804_clear" / "VLA82-017" / "replay_check_end.png",
        "图 3  VLA82-017：柜体相关操作中段，机器人与柜体无门板遮挡",
        "图 4  VLA82-017：操作末段，机器人末端和柜体操作空间可见",
    )
    add_two_images(
        doc,
        base / "video_regeneration_20260804_clear" / "VLA82-019" / "replay_check_mid.png",
        base / "video_regeneration_20260804_clear" / "VLA82-020" / "replay_check_mid.png",
        "图 5  VLA82-019：代表性柜体物体操作宽视角证据",
        "图 6  VLA82-020：代表性柜体物体操作宽视角证据",
    )
    add_image(doc, base / "video_regeneration_20260804_clear" / "VLA82-030" / "replay_check_start.png", "图 7  VLA82-030：设备/抽屉关闭操作，机器人与设备区域完整可见", width=4.7)

    heading(doc, "6. 测试结果文件说明", 1)
    para(doc, "严格 PASS 原始证据统一位于：C:\\OpenVLA-Simulator\\outputs\\midterm_testing_vla82\\full_simulation_objectwise_8x60\\expert_demos\\。每个对象目录至少包含 episode-2000.json 和 episode-2000.npz；JSON 中记录 status、predicate_success、轨迹步数和物理判定结果，NPZ 中保存图像流、动作、接触快照与谓词输出。")
    para(doc, "面向评审的视频证据位于：C:\\OpenVLA-Simulator\\outputs\\midterm_testing_vla82\\video_regeneration_20260804_wide\\ 与 C:\\OpenVLA-Simulator\\outputs\\midterm_testing_vla82\\video_regeneration_20260804_clear\\。其中宽视角回放视频的侧车 JSON 显式指向对应原严格 PASS JSON 与动作 NPZ，保证“视频呈现”与“严格物理判定”可分别核验。")

    heading(doc, "7. 阶段性结论", 1)
    para(doc, "当前测试已完成 8 类任务、60 种物体的测试组织、标注关联、场景配置和逐物体过程记录，并形成抓取/放置、容器投放、柜体相关操作、设备控制等代表性严格 PASS 结果。严格 PASS 记录显示：在已验证对象上，机器人能够基于技能阶段完成对应物理操作，且结果由真实接触、对象/设备状态变化和稳定终态共同支持。后续测试将沿用同一流程持续扩展严格 PASS 对象覆盖范围。")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
