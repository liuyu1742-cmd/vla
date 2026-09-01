from docx import Document

TARGET = r"C:\OpenVLA-Simulator\outputs\chapter5_midterm_revision\第五章（中期验收版_修订_图5.2更新版）.docx"

REPLACEMENTS = {
    87: "动态组合用于处理由多个原子操作构成的服务指令。系统先解析任务中的对象、目标区域和状态变化，再根据操作依赖关系确定执行顺序。例如，将书籍收纳至抽屉并关闭抽屉，可依次完成定位书籍、抓取书籍、打开抽屉、放入书籍、释放书籍和关闭抽屉等操作。对于将饮料瓶从冰箱放至咖啡台的任务，冰箱门开启、物品取出、目标位置放置和冰箱门复位构成连续状态变化。",
    89: "组合模块为每个原子目标选择对象、区域和动作条件均匹配的技能，并检查前一操作的输出状态是否满足后一操作的输入条件。抽屉打开是物品放入操作的前置状态，物品进入抽屉则是关闭抽屉的前置状态。通过状态衔接，技能序列能够反映对象在操作过程中的实际变化，避免仅按语言出现顺序拼接动作。",
    91: "时序约束用于描述技能之间的先后关系。单机械臂执行接触式操作时，同一时刻只安排一个末端动作；视觉观测、状态更新和路径计算等不占用末端资源的过程可与运动规划协同进行。系统在执行过程中持续维护容器开合、物体抓持、目标区域占用和物体放置等状态，使每一步操作均在满足前置条件的情况下进行。",
    93: "资源约束用于处理末端执行器、目标区域和被操作对象之间的占用关系。当两个候选操作需要同时使用同一物体、同一容器开口或同一放置区域时，系统依据对象状态和动作依赖关系调整操作顺序；当目标区域被占用或容器状态不符合要求时，则先完成相关状态转换，再继续后续操作。",
    95: "在技能执行过程中，系统以对象、容器、储物设施和目标区域为基本资源单元，记录其当前状态与可用性。对同一资源提出不相容要求的操作不同时执行；目标区域发生变化后，后续技能依据更新后的场景状态重新判断放置位置和操作路径。该过程使物品递送、整理收纳、室内摆放和储物设施开合等任务具有一致的组合逻辑。",
    96: "因此，技能检索与动态组合由候选技能选择、状态衔接、时序安排和资源协调四个环节构成。检索模块根据指令和场景条件确定对象操作，组合模块依据对象状态和操作依赖关系生成连续动作序列，从而将书籍归位、水瓶递送、餐具摆放、抽屉开合和卫浴用品取放等具体任务组织为可执行的技能流程。",
}


def replace_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


doc = Document(TARGET)
for index, text in REPLACEMENTS.items():
    replace_text(doc.paragraphs[index], text)

# Remove four display equations that introduced intermediate symbols (T, π,
# resource set and need()) without improving the technical explanation.
for index in sorted((94, 92, 90, 88), reverse=True):
    element = doc.paragraphs[index]._element
    element.getparent().remove(element)

doc.save(TARGET)
print(TARGET)
