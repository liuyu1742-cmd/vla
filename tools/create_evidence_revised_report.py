from pathlib import Path
from shutil import copy2
from copy import deepcopy
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

ROOT=Path(r'C:\OpenVLA-Simulator')
SRC=Path(r'C:\Users\sjtu101\Desktop\2. 基于视觉的人类动作捕获、解析和技能学习迁移技术研究.docx')
OUT=ROOT/'outputs'/'evidence_revised_report'
OUT.mkdir(parents=True,exist_ok=True)
DST=OUT/'基于视觉的人类动作捕获_解析和技能学习迁移技术研究_真实证据修订版.docx'
XLSX=ROOT/'outputs'/'habit_evidence_workbook'/'家居VLA_15任务_120物体_真实示教证据.xlsx'

def set_para(p,text):
    old=deepcopy(p.runs[0]._element.rPr) if p.runs and p.runs[0]._element.rPr is not None else None
    p.clear(); r=p.add_run(text)
    if old is not None: r._element.get_or_add_rPr().append(old)

def set_cell(cell,text):
    p=cell.paragraphs[0]; old=deepcopy(p.runs[0]._element.rPr) if p.runs and p.runs[0]._element.rPr is not None else None
    p.clear(); r=p.add_run(str(text))
    if old is not None:r._element.get_or_add_rPr().append(old)
    for extra in cell.paragraphs[1:]: extra.clear()

def fill_table(tbl,headers,data):
    for i,row in enumerate(tbl.rows):
        vals=headers if i==0 else (data[i-1] if i-1<len(data) else ['']*len(headers))
        for j,c in enumerate(row.cells): set_cell(c, vals[j] if j<len(vals) else '')

def after(p):
    e=OxmlElement('w:p');p._p.addnext(e);return Paragraph(e,p._parent)

copy2(SRC,DST)
doc=Document(DST)

# Read the approved 15-task/120-object ledger, rather than recreating object claims.
wb=load_workbook(XLSX,read_only=True,data_only=True)
ws=wb['任务物体清单']
objects=[]
for row in ws.iter_rows(min_row=4,values_only=True):
    if row[0]: objects.append(list(row[:8]))
if len(objects)!=120: raise RuntimeError('expected 120 objects')
groups={}
for r in objects: groups.setdefault(str(r[0]),[]).append(r)

# 4.2.5: replace unsupported performance claims with actual local validation facts.
set_para(doc.paragraphs[698],'4.2.5 实验验证')
set_para(doc.paragraphs[699],'本节采用本机已完成的数据完整性核验和端到端仿真闭环作为验证。HABIT样本已核验60个Parquet动作记录与300段RGB视频；BEHAVIOR-1K家居子集已核验1000个Parquet动作记录与1000段头戴RGB视频；DROID-100已核验100个episode、32212帧和7维动作字段。上述数据具备语言/任务文本、视觉观测与动作标签，但不同机器人本体仍需执行动作映射后才可共同训练。')
set_para(doc.paragraphs[700],'表4-1给出本阶段可复核的实验性结果。此处不以未完成的实体机器人实验替代数据证据：表中“可进入转换”表示已具备真实示教三要素并已在本机落盘；“待下载映射”表示公开数据集存在真实示教，但本项目尚未将对应episode下载并映射为目标7维动作。')
set_para(doc.paragraphs[701],'表 4-1 本地数据与动作接口验证结果')
fill_table(doc.tables[8],['验证对象','本地规模','视觉/动作字段','结论','限制'],[
['HABIT样本','60 Parquet；300 RGB MP4','5路RGB；FR3双臂动作','可进入转换','需双臂到目标臂映射'],
['BEHAVIOR-1K子集','1000 Parquet；1000 RGB MP4','头戴RGB；R1Pro 23D动作','可进入转换','需23D到7D映射'],
['DROID-100','100 episode；32212帧','多RGB；7D action','可进入转换','需LeRobot/OXE字段适配'],
['OpenVLA推理服务','CUDA本地启动','图像+指令到7维动作','服务可用','不等同于任务成功率'],
['Humanoid Everyday','246项本地任务元数据','公开RGB/动作数据','待选集下载','尚未本地配对']])

# 4.4.4: be explicit that no physical Sim-to-Real experiment has been conducted.
set_para(doc.paragraphs[767],'4.4.4 虚实迁移实验结果')
set_para(doc.paragraphs[768],'本项目当前仅完成RoboCasa仿真闭环与跨进程动作执行验证，尚未接入实体相机、真实机械臂或力传感器。因此不报告“真实成功率”、Sim-to-Real gap或领域随机化提升百分比。后续实体部署前应以同一任务、同一成功判据分别在仿真和真实平台各运行不少于20个随机种子，再计算成功率与置信区间。')
set_para(doc.paragraphs[769],'表4-4列出已完成的可复核验证及尚未执行的实体迁移项目。该写法将“仿真可运行”与“已经完成虚实迁移”严格区分，避免由仿真结果外推实体性能。')
set_para(doc.paragraphs[770],'当前可确认的结论是：OpenVLA服务、RoboCasa观察—动作IPC及仿真动作执行已打通；实体机器人、相机标定、控制频率、接触安全和物理参数辨识仍是下一阶段的必要条件。')
set_para(doc.paragraphs[771],'表 4-4 当前虚实迁移验证状态')
fill_table(doc.tables[11],['验证项目','本地结果','证据','结论'],[
['RoboCasa观察—动作闭环','已完成','OpenVLA TCP + RoboCasa执行日志','仿真接口可用'],
['组织玩具入柜仿真任务','3/3 held-out seed成功（混合闭环）','seed 101/102/103验收报告','仅仿真任务证据'],
['实体相机与机械臂','未接入','无实体设备记录','不可报告真实成功率'],
['Sim-to-Real gap','未测量','无同任务实体对照','后续按统一协议测量']])

# 4.6: use the one real held-out evaluation, making hybrid assistance explicit.
set_para(doc.paragraphs[807],'4.6 实验验证')
set_para(doc.paragraphs[808],'本节报告本项目已经完成的OpenVLA本地训练与仿真验收。实验对象为RoboCasa中的“put away the toy in the cabinet”任务；训练使用Dagger-r3 LoRA适配器，评估种子101、102、103与训练集严格隔离。结果是混合闭环结果，不应表述为纯自主VLA性能。')
set_para(doc.paragraphs[809],'4.6.1 OpenVLA本地微调与数据隔离')
set_para(doc.paragraphs[810],'Dagger-r3在augmented-r2适配器基础上训练1个epoch，计划更新数与完成更新数均为19261。训练清单的held-out seeds为101、102、103，与训练种子无重叠。该结果验证了本地微调流水线与数据隔离，而非跨平台泛化结论。')
set_para(doc.paragraphs[811],'表4-7记录可复核的训练配置；不保留与本项目无关的LIBERO全参数微调成功率。')
set_para(doc.paragraphs[812],'表 4-7 本地OpenVLA适配器训练与数据隔离')
fill_table(doc.tables[14],['项目','本地结果','证据','解释'],[
['初始适配器','augmented-r2','训练报告','作为Dagger-r3初始化'],
['训练轮数','1 epoch','训练报告','完成一次增量训练'],
['更新次数','19261 / 19261','训练报告','计划与完成一致'],
['held-out','101、102、103','训练清单','与训练集零重叠']])
set_para(doc.paragraphs[814],'4.6.2 保留种子仿真闭环结果')
set_para(doc.paragraphs[815],'在每个保留种子上，最大决策步数设置为300、动作重复为2。3个种子均满足仿真器成功、物体进入储物区和夹爪释放三个原生成功判据。平均每个种子执行约212.3个决策步，平均单次模型推理约0.53秒（seed 101记录）。')
set_para(doc.paragraphs[816],'表4-8给出保留种子结果。它证明该特定RoboCasa任务在当前混合闭环中可完成；不能推出15类任务、120种物体均已训练成功，也不能等价为纯OpenVLA自主成功率。')
set_para(doc.paragraphs[817],'为保证接触阶段稳定性，运行器使用状态感知定位恢复、校准抓取时序和最终接触伺服。这些模块是系统工程约束，不是模型端到端学习出的行为，应在验收时单独披露。')
set_para(doc.paragraphs[818],'表 4-8 保留种子RoboCasa混合闭环验收结果')
fill_table(doc.tables[15],['保留种子','是否成功','决策步数','仿真步数'],[['101','是','211','421'],['102','是','214','427'],['103','是','212','423'],['汇总','3/3（100%）','637','1271'],['判据','入柜且释放','held-out','非纯自主VLA']])
set_para(doc.paragraphs[820],'4.6.3 运行模式与案例说明')
set_para(doc.paragraphs[821],'本项目未部署ALOHA或实体双臂机器人，故不报告ALOHA性能。表4-9统计当前3个保留种子中的执行模式：OpenVLA负责受保护的动作建议，但定位恢复、校准抓取和最终接触阶段含有规则/安全控制。')
set_para(doc.paragraphs[822],'表4-9与案例图用于说明系统实际组成。最终接触伺服占300个决策步骤；因此本节的3/3成功仅计为“混合闭环仿真验收”，不计入纯自主VLA验收。')
set_para(doc.paragraphs[823],'表 4-9 混合闭环执行模式汇总')
fill_table(doc.tables[16],['执行模式','决策步数','作用','归属'],[['OpenVLA guarded','160','视觉语言动作建议','模型'],['状态感知定位恢复','46','偏差状态恢复','安全/专家'],['校准抓取时序','131','抓取窗口控制','规则控制'],['最终接触伺服','300','入柜与释放稳定','规则控制'],['合计','637','3个held-out种子','混合闭环']])
set_para(doc.paragraphs[824],'案例图4-4来自seed 101的实际RoboCasa rollout视频，展示从接近目标到物体进入柜体并释放的三个时刻。该图仅作为仿真执行轨迹证据，不代表实体机器人演示。')
set_para(doc.paragraphs[826],'本章的本地证据表明：OpenVLA推理服务、真实示教数据的转换入口及RoboCasa闭环已建立；Dagger-r3在一个保留的整理入柜任务上完成3/3混合闭环仿真验收。尚未完成的部分包括实体相机/机械臂验证、纯自主VLA长回合成功率，以及15类任务、120种物体的逐项训练与评测。因此后续应先完成数据本体与动作空间映射，再按任务类别开展受控微调和实体安全验证。')

# Chapter 5: map exactly to the approved 15 categories; keep original local heading/body slots.
task_names=['家庭物品整理与收纳','室内物品取用、递送与摆放','室内垃圾分类与投放','餐具整理与餐后清洁','食品存放与食材管理','烹饪、加热与饮品服务','家电与水电设施控制','室内卫生清洁与物件维护','衣物清洗与织物护理','家庭维修与工具使用','室内布置与家具陈设','浴室用品与个人卫生服务','工作学习区域服务','居家休闲与娱乐用品服务','室内出入与随身物品服务']
task_desc=['面向书籍、文件夹、笔袋、笔记本、签字笔、柜子和收纳箱等对象，执行分类、放入、叠放、开关柜门和归位。','面向水瓶、托盘、容器、瓶子和室内台面等对象，执行取用、承托、递送和目标摆放。','面向纸巾、废纸、饮料罐、垃圾桶、回收篮和厨余桶等对象，执行识别、分类和投放。','面向盘子、碗、叉子、玻璃杯、锅、锅盖、水槽和洗碗机，执行收集、清洗、装载和归位。','面向冰箱、保鲜盒、厨房柜、密封罐、干粮收纳盒和食用油瓶，执行开关、存放、密封和分类。','面向平底锅、炉灶、烤箱、微波炉、烤盘、餐食食材和微波食品盒，执行放置、加热、取出和端送。','面向灯开关、水龙头、电水壶、烤面包机、插线板和微波炉门，执行按压、旋转、开关和档位控制。','面向掸子、刷子、海绵、清洁布、桌面、白板、灯罩和置物架，执行除尘、擦拭、刷洗和复位。','面向洗衣机、枕套、帽子、毛巾、T恤、牛仔裤、内衣和洗衣篮，执行装载、清洗、折叠和收纳。','面向工具箱、电钻、钳子、手电筒、扳手、螺丝刀、锤子和万用表，执行取放、对准、旋转和检测。','面向海报、墙钉、椅子、沙发、床头柜、花瓶、相框和手机支架，执行对齐、固定、推拉、摆正和陈设。','面向洗涤剂、卫生巾盒、皂液器、牙膏、牙刷、洗发水、卫生纸和洗手液，执行取用、补充与归位。','面向显示器、键盘、鼠标、订书机、笔记本电脑、打孔器、计算器和摄像头，执行摆正、按键、清洁与角度调整。','面向数码相机、三脚架、棋子、游戏手柄、运动器材、桌游、拼图和耳机，执行取放、摆放、操作和收纳。','面向鞋子、外套、雨伞、随身包、拉链头、钥匙串、门把手和鞋架，执行挂放、开合、归位与门口整理。']
set_para(doc.paragraphs[828],'5.1 技能库总体设计')
set_para(doc.paragraphs[829],'本项目技能库以“任务类别—物体—对象操作—示教证据—训练状态”为核心数据模型。每条技能记录绑定自然语言任务文本、RGB视频/视觉观测、动作字段与机器人本体信息；只有具备真实示教证据的记录进入候选训练集。')
set_para(doc.paragraphs[830],'技能执行层将高层任务分解为接近、抓取/接触、移动、放置/释放、开关/旋转、清洁等动作原语，并将动作原语参数化为目标物体、目标区域、夹爪状态和安全约束。跨数据集训练前统一转换视觉字段、语言字段和目标7维动作接口。')
set_para(doc.paragraphs[831],'当前证据清单包含15类室内日常任务和120种具体物体。HABIT、BEHAVIOR-1K和DROID-100的已本地配对数据可进入转换流水线；Humanoid Everyday与Dynamic Intelligence的对象仅标为“公开真实示教、待选集下载与动作映射”，不计为已完成训练。')
set_para(doc.paragraphs[833],'图 5.1 技能库总体框架图（任务—物体—操作—示教证据—训练状态）')
set_para(doc.paragraphs[834],'5.2 技能分类体系')
set_para(doc.paragraphs[835],'分类以实际可获得的室内机器人示教为约束，而非以概念性智能家居功能凑类。15类任务边界按主要操作目的划分；每类当前列出8种物体，共120种，避免把户外、采购、儿童专属或重复整理类任务纳入范围。')
for i,(name,desc) in enumerate(zip(task_names,task_desc)):
    h=836+i*3; set_para(doc.paragraphs[h],f'5.2.{i+1} {name}'); set_para(doc.paragraphs[h+1],desc); set_para(doc.paragraphs[h+2],'')
set_para(doc.paragraphs[880],'技能分类体系、物体、对象操作及数据来源详见表5-1至表5-9；每个物体均保留对应任务文本和训练转换状态。')
set_para(doc.paragraphs[881],'表 5-1 真实示教约束下的15类室内任务总表')
summary=[]
for i,n in enumerate(task_names): summary.append([f'{i+1:02d}',n,'8',C if False else '见表5-2至表5-9','真实示教/待下载映射'])
fill_table(doc.tables[17],['编号','任务类别','物体数','代表物体','证据状态'],summary)
set_para(doc.paragraphs[883],'5.3 物体操作技能详细列表')
set_para(doc.paragraphs[884],'表5-2至表5-9按任务类别列出120种对象、其具体操作和训练状态。数据状态为“本地真实配对”的条目可进入动作转换；“公开真实示教”的条目需先完成所需episode下载与动作映射。')
caps=['表 5-2 任务1—2：整理、递送与摆放','表 5-3 任务3—4：垃圾与餐后清洁','表 5-4 任务5—6：食品管理与加热服务','表 5-5 任务7—8：设施控制与卫生清洁','表 5-6 任务9—10：织物护理与维修','表 5-7 任务11—12：室内陈设与卫生用品','表 5-8 任务13—14：工作学习与休闲用品','表 5-9 任务15：出入与随身物品']
for pidx,cap in zip([885,887,889,891,893,895,897,899],caps):set_para(doc.paragraphs[pidx],cap)
for ti,pairs in zip(range(18,26),[(0,1),(2,3),(4,5),(6,7),(8,9),(10,11),(12,13),(14,)]):
    data=[]
    for ci in pairs:
        for r in groups[f'{ci+1:02d}']:
            data.append([task_names[ci],r[3],r[4],r[7]])
    fill_table(doc.tables[ti],['任务类别','物体','对象操作','训练状态'],data)
set_para(doc.paragraphs[922],'5.5 本章小结')
set_para(doc.paragraphs[923],'本章按真实示教证据构建了15类室内日常任务、120种物体及其对象操作清单。技能库以任务、物体、操作、示教来源和训练状态为可追溯单元；它支持后续按数据本体、动作空间与目标机械臂逐项转换和评测。当前表中的“公开真实示教”并不等价于已训练成功，后续需完成选集下载、动作映射、微调与统一成功判据评估。')

# Insert an actual rollout montage only if OpenCV can decode the recorded rollout.
try:
    import cv2, numpy as np
    video=ROOT/'outputs'/'formal_skill_dagger_r3_hybrid_v3_eval'/'seed_101'/'rollout.mp4'
    cap=cv2.VideoCapture(str(video)); n=int(cap.get(cv2.CAP_PROP_FRAME_COUNT)); frames=[]
    for k in [0,max(0,n//2),max(0,n-1)]:
        cap.set(cv2.CAP_PROP_POS_FRAMES,k); ok,f=cap.read()
        if ok: frames.append(cv2.cvtColor(f,cv2.COLOR_BGR2RGB))
    cap.release()
    if len(frames)==3:
        h=min(x.shape[0] for x in frames); frames=[cv2.resize(x,(int(x.shape[1]*h/x.shape[0]),h)) for x in frames]
        montage=np.concatenate(frames,axis=1); fig=OUT/'figure_4_4_rollout.png'; cv2.imwrite(str(fig),cv2.cvtColor(montage,cv2.COLOR_RGB2BGR))
        pic=after(doc.paragraphs[824]); pic.alignment=WD_ALIGN_PARAGRAPH.CENTER; pic.add_run().add_picture(str(fig),width=Inches(6.0))
        cap_p=after(pic); cap_p.style=doc.paragraphs[697].style; cap_p.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap_p.add_run('图 4-4 seed 101：RoboCasa整理入柜任务的实际仿真rollout关键帧')
except Exception as e:
    print('figure skipped',e)

doc.save(DST)
print(DST)
