from pathlib import Path
from copy import deepcopy
import json
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOC=Path(r"outputs/chapter5_midterm_revision/第五章（中期验收版_修订_分类与公式版）.docx")
AUDIT=Path(r"outputs/chapter5_midterm_revision/chapter5_object_alignment_after_public_polish.json")

def operation(task,obj,instruction):
    if task=="5.2.5":
        return {"垃圾桶":"多罐汽水投放：客厅→厨房垃圾桶","饮料罐":"多罐汽水投放：客厅→厨房垃圾桶","海绵":"海绵取放：台面→橱柜","清洁刷":"洗碗刷取放：橱柜→台面","清洁喷雾瓶":"喷雾瓶取放：橱柜→台面"}.get(obj,"按源指令完成清洁用品操作")
    if task=="5.2.6":
        if obj in {"削皮刀","砧板","餐刀"}: return "洋葱切丁：水槽→砧板→碗，刀具与砧板回放水槽"
        if obj=="平底锅": return "卷心菜和辣椒切丁并在炉灶平底锅中烹调"
        if obj=="混合碗": return "按蔬菜类别分入三个混合碗"
        if obj=="炉灶": return "导航至厨房炉灶"
        if obj=="烤箱": return "关闭烤箱门"
        if obj=="烤面包机下层烤盘": return "将烤箱下层烤盘完全推入"
        if obj=="烹饪托盘": return "将烤箱上层托盘完全拉出"
        if obj=="玻璃罐": return "开柜、开罐、放入熟香肠、关罐并归回"
        if obj=="碗": return "苹果：平底锅→碗"
        if obj in {"打蛋器","擀面杖","茶壶","调味罐","量杯","锅"}: return {"打蛋器":"台面→抽屉","擀面杖":"台面→抽屉","茶壶":"台面→橱柜","调味罐":"台面→橱柜","量杯":"台面→抽屉","锅":"台面→橱柜"}[obj]+"单件归位"
    if task=="5.2.7": return f"{obj}：工具台面→工具箱，保持工具箱在台面并合上箱盖"
    if task=="5.2.8":
        return {"书籍":"书籍：床上→床头柜；两只凉鞋并排放在床边","水瓶":"水瓶：橱柜→台面","盒装饮料":"盒装饮料：橱柜→台面","罐装物":"罐装物：橱柜→台面","饮料瓶":"两瓶饮料：厨房冰箱→客厅咖啡台，完成后关冰箱","马克杯":"马克杯：橱柜→台面"}.get(obj,"按源指令完成物品递送")
    if task=="5.2.9":
        return {"密封保鲜盒":"密封保鲜盒：台面→橱柜","床头柜":"书籍：床→床头柜，凉鞋并排放置","收纳盒":"盒装饮料：台面→橱柜","收纳箱":"六本书：书架→客厅地面收纳箱","收纳篮":"卧室篮中物品分类归位至浴室指定区域","纸箱":"六本书：书架→客厅地面纸箱"}.get(obj,"按源指令完成整理收纳")
    if task=="5.2.10":
        return {"冰箱层架":"苹果：冰箱抽屉→冰箱层架","冰箱抽屉":"完全关闭冰箱抽屉","冰箱门":"关闭冰箱门","微波炉门":"关闭微波炉门","抽屉":"关闭左侧抽屉","橱柜门":"关闭橱柜门","洗碗机上层架":"将洗碗机上层架完全推入","洗碗机门":"关闭洗碗机门","烤箱门":"关闭烤箱门","烤面包机烤箱门":"打开烤面包机烤箱门"}.get(obj,"按源指令完成设施开合")
    if task=="5.2.11":
        return {"墙钉":"海报悬挂到厨房墙钉","数码相机":"数码相机安装到卧室三脚架","海报":"厨房台面海报→墙钉悬挂","相机三脚架":"接收并支撑数码相机","蜡烛":"蜡烛：橱柜→台面"}.get(obj,"按源指令完成室内摆放或安装")
    if task=="5.2.12":
        routes={"咖啡杯":"橱柜→台面","夹子":"橱柜→台面","披萨刀":"橱柜→台面","木勺":"台面→橱柜","水罐":"台面→橱柜","汤勺":"台面→橱柜","洗碗刷":"台面→橱柜","滤盆":"台面→橱柜","盘子":"苹果：平底锅→盘子","铝箔纸":"台面→橱柜"}; return routes.get(obj,"按源指令完成餐具摆放")
    if task=="5.2.13":
        return {"棒球帽":"两顶棒球帽：台面→洗衣机清洗","毛巾":"毛巾分入两个篮子","洗衣机":"运行洗衣机清洗两顶棒球帽","鞋子":"两双运动鞋和两双凉鞋：走廊地面→鞋架并按类别并排","鞋架":"接收鞋类并保持运动鞋、凉鞋分别并排"}.get(obj,"按源指令完成衣物鞋类服务")
    if task=="5.2.14":
        if obj in {"办公椅","文件夹","显示器","电脑","笔记本","键盘","鼠标"}: return {"办公椅":"移至书桌旁","文件夹":"转椅→书桌并置于鼠标旁","显示器":"放在书桌上","电脑":"放在书桌下","笔记本":"叠放在文件夹上","键盘":"放在显示器旁","鼠标":"放在键盘旁"}[obj]
        return {"笔盒":"铅笔和两支笔→笔盒，笔盒留在桌面","笔记本电脑":"床→书桌并合上电脑","订书机":"书柜→桌面","铅笔":"放入笔盒","显示器":"放在书桌上"}.get(obj,"按源指令完成工作区摆放")
    if task=="5.2.15":
        return {"卫浴搁板":"接收卫生巾盒并保持浴室用品分区","卫生巾盒":"卧室篮→浴室搁板","固体香皂":"台面→橱柜","洗涤剂":"卧室篮→洗手池下方并排放置","牙刷":"放入洗手台上的杯中","牙膏":"放入洗手台上的杯中","皂液器":"橱柜→台面"}.get(obj,"按源指令完成卫浴用品归位")
    return "按源指令完成对象操作"

def move_before(parent,element,new): parent.insert(parent.index(element),new)

def main():
    d=Document(DOC); records=json.loads(AUDIT.read_text(encoding='utf-8'))['records']
    groups={}
    for r in records: groups.setdefault(r['task_code'],[]).append(r)
    old=d.tables[-1]; parent=old._tbl.getparent(); old_index=parent.index(old._tbl)
    # remove old monolithic table and its immediately preceding Table 5-8 caption
    parent.remove(old._tbl)
    for p in list(d.element.body.iterchildren()):
        if p.tag.endswith('}p') and p.text if False else False: pass
    for p in list(d.paragraphs):
        if p.text.strip().startswith('表5-8 已核验VLA物体操作样本'):
            p._element.getparent().remove(p._element)
    alln=0
    task_names={
      '5.2.5':'室内卫生清洁','5.2.6':'烹饪与加热辅助','5.2.7':'设备维护与工具使用','5.2.8':'物品递送','5.2.9':'整理收纳','5.2.10':'储物设施开合','5.2.11':'室内安装与布置','5.2.12':'餐具与容器摆放','5.2.13':'衣物鞋类与玄关归位','5.2.14':'工作学习区服务','5.2.15':'卫浴用品与个人卫生服务'}
    anchor=next(x for x in d.element.body.iterchildren() if x.tag.endswith('}p') and '5.5' in ''.join(x.itertext()))
    for idx,task in enumerate(sorted(groups,key=lambda x: tuple(int(v) for v in x.split('.'))),1):
        cap=d.add_paragraph(f"表5-8-{idx} {task_names[task]}对象操作技能（{len(groups[task])}种）"); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        table=d.add_table(rows=1,cols=4); table.style=d.tables[1].style
        for c,v in zip(table.rows[0].cells,['序号','技能名称','操作对象','操作类型']): c.text=v
        for r in groups[task]:
            alln+=1; row=table.add_row(); row.cells[0].text=str(alln); row.cells[1].text=task_names[task]; row.cells[2].text=r['object']; row.cells[3].text=operation(task,r['object'],r['instruction'])
        move_before(d.element.body,anchor,cap._p); move_before(d.element.body,anchor,table._tbl)
    d.save(DOC); print('tables_created',len(groups),'objects',alln)
if __name__=='__main__': main()


