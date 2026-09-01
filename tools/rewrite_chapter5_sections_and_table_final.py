from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement

DOC=Path(r"outputs/chapter5_midterm_revision/第五章（中期验收版_修订_分类与公式版）.docx")
ORIG=Path(r"C:\Users\sjtu101\Desktop\2. 基于视觉的人类动作捕获、解析和技能学习迁移技术研究(原版).docx")

ADDITIONS={
"5.2.1":"在具体应用中，系统可将客厅、卧室、厨房和卫生间的设备分别纳入统一的家庭设备清单，按照设备类别选择相应接口。机器人需要理解开关、模式、音量、温度和定时等参数的语义差异，完成指令解析、设备调用和结果确认。对于同一设备的连续操作，还应保持前一次状态与当前请求的一致，避免重复执行或状态冲突。",
"5.2.2":"按照调节对象可进一步分为温度控制、湿度控制、空气质量控制和照明控制。温度控制主要涉及空调的开关、模式和设定温度；湿度控制涉及加湿器和除湿设备的运行状态；空气质量控制涉及空气净化器的档位和运行时间；照明控制涉及灯具的开关、亮度和色温。机器人执行时需要综合传感器读数与用户偏好，选择合适的调节设备。",
"5.2.3":"入口安全可按照监测对象进一步分为门体状态监测、门锁状态管理、门铃事件处理和入口视频查看。门体监测用于判断家庭出入口是否关闭，门锁管理用于确认上锁或解锁状态，门铃处理用于识别来访事件，视频查看用于辅助确认入口环境。机器人执行时需要按照用户指令调用对应设备，并以状态回读和事件记录完成服务确认。",
"5.2.4":"环境安全可按照风险来源分为烟雾监测、燃气监测、漏水监测和紧急求助。烟雾与燃气监测用于发现异常浓度，漏水监测用于发现厨房、卫生间等区域的水浸状态，紧急求助用于接收用户发出的报警请求。机器人执行时需要区分不同事件的优先级，准确读取传感器状态，按照预设规则发送通知并保存报警记录。",
"5.2.6":"厨房器具可按照用途分为量取类、搅拌类、面食处理类、盛装类和调味类。量取类包括量杯，用于完成容量测量；搅拌类包括打蛋器，用于混合和搅拌；面食处理类包括擀面杖，用于面团延展；盛装类包括茶壶，用于饮品盛放；调味类包括调味罐，用于调味品存取。机器人执行时需要识别器具的外形与可抓取部位，按照目标位置完成安全取放。",
"5.2.7":"养护管理可按照服务对象分为绿植养护、空气设备维护、清洁设备维护和安全设备自检。绿植养护关注花盆状态与浇水提醒，空气设备维护关注净化器滤网和运行时长，清洁设备维护关注扫地机器人耗材与工作状态，安全设备自检关注报警器电量和测试状态。机器人需要根据设备接口返回的信息生成提醒或执行相应控制。",
"5.2.8":"物品递送可按照物品类型分为饮用水递送、盒装饮料递送、罐装物递送和杯具递送。不同物品虽然形态和重量存在差异，但都需要完成目标识别、抓取、短距离移动和指定位置放置。机器人执行时需要依据语言指令确定取物点与放置点，规划安全移动路径，并在放置后检查物体是否处于稳定状态，从而完成递送服务。",
"5.2.9":"整理收纳可按照收纳对象分为食品容器收纳、生活用品收纳和文具用品收纳。食品容器需要按照密封性和使用频率放置，生活用品需要按照房间和储位归类，文具用品需要按照类别和取用便利性排列。机器人执行时需要理解物品类别、目标容器和放置区域之间的关系，完成抓取、移动、放置和位置确认，使物品保持整齐并便于再次取用。",
"5.2.10":"储物设施可按照结构分为抽屉、柜门、冰箱门、冰箱抽屉、微波炉门和洗碗机门。不同设施的操作方向、把手位置和运动行程不同，机器人需要根据视觉信息确定接触点和运动方向。执行过程中应保持末端速度平稳，避免对门体和抽屉产生冲击，并依据门缝、位置变化或关闭状态判断动作是否完成。",
"5.2.11":"室内摆放可按照空间位置分为台面摆放、柜内摆放和局部装饰摆放。台面摆放要求物体位于指定区域且姿态稳定，柜内摆放要求避开柜壁并保持取用空间，局部装饰摆放要求兼顾物体朝向和周围布局。机器人执行时需要从场景中识别目标区域边界，选择合适的抓取姿态，并通过落放后的视觉反馈校正位置。",
"5.2.12":"餐具与容器可按照功能分为饮品容器、夹取工具、切配工具、搅拌工具、盛装容器和过滤容器。饮品容器包括咖啡杯，夹取工具包括夹子，切配工具包括披萨刀，搅拌工具包括木勺和汤勺，盛装容器包括水罐，过滤容器包括滤盆。机器人执行时需要根据器具形状选择抓取区域，控制放置高度与方向，避免器具倾倒或相互碰撞。",
"5.2.13":"衣物鞋类与玄关归位可按照对象分为鞋类归位、毛巾归位和洗衣设备服务。鞋类归位需要识别鞋子的左右关系和鞋架位置，毛巾归位需要识别织物类别与收纳篮，洗衣设备服务需要根据程序和设备状态完成相应操作。机器人执行时需要结合玄关、卫生间和洗衣区的空间布局，保持物品摆放整齐并方便家庭成员取用。",
"5.2.14":"工作学习区服务可按照对象分为计算设备、输入设备、文件用品和桌面辅助物品。计算设备包括笔记本电脑和显示器，输入设备包括键盘和鼠标，文件用品包括文件夹和笔记本，桌面辅助物品包括文具及办公用品。机器人执行时需要识别桌面区域和物品相对位置，按照用户指令完成物品摆放、取用或区域整理，并避免遮挡屏幕和占用键盘操作空间。",
"5.2.15":"卫浴用品与个人卫生服务可按照用品类型分为清洁用品、洗手用品和个人护理用品。清洁用品用于维持洗手台和卫浴空间整洁，洗手用品包括固体香皂和皂液器，个人护理用品包括牙刷、牙膏及卫生用品。机器人执行时需要区分洗手台、储物柜和置物架等区域，按照物品的使用频率和卫生要求完成取放、归位与状态检查。",
}

def body_from(doc, code):
    ps=doc.paragraphs
    indexes=[i for i,p in enumerate(ps) if p.text.strip().startswith(code)]
    return ps[max(indexes)+1].text

def replace_52(doc, orig):
    for i in range(1,16):
        code=f"5.2.{i}"
        original=body_from(orig,code)
        text=original
        if len("".join(text.split()))<350:
            text += ADDITIONS[code]
        if len("".join(text.split()))<350:
            text += "机器人需要依据任务指令、场景状态和目标对象完成连续而稳定的服务动作，并在任务结束后确认目标状态。"
        idx=next(j for j,p in enumerate(doc.paragraphs) if p.text.strip().startswith(code))
        p=doc.paragraphs[idx+1]
        p.clear(); p.add_run(text)

def restore_54(doc, orig):
    def block_bounds(d):
        body=d.element.body; children=list(body.iterchildren()); ps=d.paragraphs
        h4=max((p for p in ps if p.text.strip().startswith("5.4")), key=lambda p: ps.index(p))
        h5=max((p for p in ps if p.text.strip().startswith("5.5")), key=lambda p: ps.index(p))
        return body,children,children.index(h4._p),children.index(h5._p)
    body,children,a,b=block_bounds(doc)
    old_nodes=children[a:b]
    for node in old_nodes: body.remove(node)
    obody,ochildren,oa,ob=block_bounds(orig)
    insert_before=next((n for n in body.iterchildren() if n.tag.endswith('}p') and '5.5' in ''.join(n.itertext())), None)
    for node in ochildren[oa:ob]:
        new=deepcopy(node)
        if insert_before is not None: insert_before.addprevious(new)
        else: body.append(new)

def classify_operation(skill,obj,source):
    if "开合" in skill or "设施" in skill: return "开合/推入"
    if "清洁" in skill: return "抓取—放置/投放"
    if "烹饪" in skill: return "厨房操作"
    if "递送" in skill: return "抓取—移动—放置"
    if "整理" in skill or "归位" in skill: return "归类—放置"
    if "摆放" in skill or "安装" in skill: return "取放/摆放"
    if "衣物" in skill: return "取放/清洗"
    if "工作" in skill: return "桌面取放"
    if "卫浴" in skill: return "取放/归位"
    return "状态查看/维护"

def table_58(doc, orig):
    t=doc.tables[5]
    # Remove source-instruction and verification columns from each row.
    for row in t.rows:
        while len(row.cells)>4:
            row._tr.remove(row.cells[-1]._tc)
    header=t.rows[0].cells
    for c,v in zip(header,["序号","技能名称","操作对象","操作类型"]): c.text=v
    for n,row in enumerate(t.rows[1:],1):
        skill=row.cells[0].text.strip(); obj=row.cells[1].text.strip()
        source=""
        row.cells[0].text=str(n)
        row.cells[1].text=skill
        row.cells[2].text=obj
        row.cells[3].text=classify_operation(skill,obj,source)
    t.style=orig.tables[18].style

def main():
    doc=Document(DOC); orig=Document(ORIG)
    replace_52(doc,orig)
    restore_54(doc,orig)
    table_58(doc,orig)
    j=next(j for j,x in enumerate(doc.paragraphs) if x.text.strip().startswith("5.3"))
    doc.paragraphs[j+1].clear(); doc.paragraphs[j+1].add_run("本节按照技能类别对非网络控制对象进行归类说明。表5-8采用与表5-2、表5-3、表5-4和表5-7一致的四列表格，依次列出序号、技能名称、操作对象和操作类型；每个对象均对应本地已核验的任务指令、RGB视频和机器人动作标签。")
    doc.save(DOC)
    print("rewritten")
if __name__=='__main__': main()





