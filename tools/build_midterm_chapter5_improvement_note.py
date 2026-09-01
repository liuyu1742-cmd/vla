"""Build a standalone, evidence-aware Chapter 5 midterm improvement note."""
from __future__ import annotations

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pathlib import Path


OUT = Path("outputs/chapter5_midterm_improvement/第五章任务体系改进与中期验收实施说明.docx")


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), fill)
    tc_pr.append(item)


def set_cell_text(cell, value: str, bold: bool = False):
    cell.text = ""
    run = cell.paragraphs[0].add_run(value)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True)
        shade(table.rows[0].cells[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


def add_para(doc, text, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(5)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)
    return p


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.4); sec.bottom_margin = Cm(2.4)
    sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"; style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    r = title.add_run("第五章任务体系改进与中期验收实施说明")
    r.bold = True; r.font.name = "Microsoft YaHei"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei"); r.font.size = Pt(18)
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("依据原版、修改版文档及当前 OpenVLA 项目数据与代码状态编制")
    sr.font.name = "Microsoft YaHei"; sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei"); sr.font.size = Pt(10.5); sr.font.color.rgb = RGBColor(89,89,89)

    add_heading(doc, "一、编制目的与边界", 1)
    add_para(doc, "本说明仅提出第五章后续修改和中期验收的改进方案，不修改“原版”或“修改”文档。核心原则是将任务定义、数据证据、代码接口和验收指标逐项对应：凡需要机械臂执行的任务，必须具备语言指令、RGB观测和机器人动作标签，并在仿真闭环中完成验证；凡通过智慧家居应用或局域网接口完成的任务，必须单列为网络控制型技能，不得以视觉模仿学习效果替代其验收。")
    add_para(doc, "当前项目已具备 OpenVLA 本地推理、RoboCasa 仿真闭环、BridgeData V2、BEHAVIOR-1K、HABIT 和 RoboCasa365 等数据来源。其中，RoboCasa365 已核验 7356 个回合均含官方任务指令、三路 RGB 视频和 12 维机器人动作；BEHAVIOR-1K 与 HABIT 的本地样本也已建立任务级视频—动作配对证据。上述条件可支撑中期阶段的单物体、单状态变化任务，但不足以直接宣称复杂烹饪、长序列收纳或跨房间递送已完成学习。")

    add_heading(doc, "二、第五章15类任务的难度重评与取舍原则", 1)
    add_para(doc, "原版第5.2.1至5.2.4及第5.2.10至5.2.14的若干任务以家电、环境、安防、能源和场景联动为核心，其主要执行链路是“语音或界面指令—设备状态读取—网络控制命令—状态回读”。这类任务不依赖机械臂抓取，也不要求从视频学习末端轨迹，因而适合保留为中期验收中的低风险软件控制模块；但其数据、代码和指标应来自长虹智慧家居 APP 的授权接口、局域网协议或可复现实验桩，而不是 VLA 数据集。")
    add_para(doc, "修改版中食品管理、烹饪服务、复杂整理收纳等任务包含多对象识别、容器开合、顺序依赖、精确放置甚至加热状态判断。它们在后期可作为扩展方向，但不宜列入中期的核心验收任务。中期应优先选择目标对象唯一、目标位置明确、动作原语少、失败可自动判定的任务，例如按钮按压、柜门开合、杯子放置、单件物品入柜和固定路线递送。")
    rows = [
        ("网络控制型", "5.2.1 家电综合管理", "保留", "应用/网络命令控制灯、空调、电视等；以命令成功率、状态回读一致率验收", "★☆☆☆☆"),
        ("网络控制型", "5.2.2 环境调节", "保留", "读取温湿度等传感器并下发阈值控制；不要求机械臂", "★☆☆☆☆"),
        ("网络控制型", "5.2.3、5.2.4 智慧安防", "保留", "入口/环境状态告警、联动与日志；需接口或仿真桩证据", "★★☆☆☆"),
        ("VLA优先", "5.2.8 物品递送", "保留并收缩", "仅保留单物体、固定起终点递送；不做跨房间自主导航", "★★☆☆☆"),
        ("VLA优先", "5.2.9 整理收纳", "保留并收缩", "仅保留单件归位或单容器放置；删除多物体排序/堆叠", "★★★☆☆"),
        ("VLA后期", "5.2.5 清洁、5.2.6 烹饪、5.2.7 养护", "后移", "接触、持续轨迹、状态识别和多步骤依赖较强", "★★★★☆"),
        ("网络控制型", "5.2.10、5.2.11、5.2.12", "保留", "能源、设施监控和场景联动均应以状态机/接口测试实现", "★☆☆☆☆"),
        ("网络控制型", "5.2.13、5.2.14", "保留为演示", "健康提醒、娱乐设备控制需明确可用设备接口；无接口时只保留仿真演示", "★★☆☆☆"),
        ("待界定", "5.2.15 其他综合任务", "拆分或删除", "不应作为无法验收的兜底类别", "—"),
    ]
    add_table(doc, ["任务类型", "原版章节", "建议", "中期验收边界", "难度"], rows, [2.0, 3.0, 2.2, 7.4, 1.4])

    add_heading(doc, "三、中期阶段建议保留的最小可验收任务集", 1)
    add_para(doc, "为降低中期风险，建议不以“15类全覆盖”作为本阶段通过条件，而采用“网络控制任务＋4个VLA单物体闭环”双轨验收。网络控制任务沿用原版第5.2.1、5.2.2、5.2.3、5.2.4以及原表中5.2、5.3、5.4、5.7对应条目，但必须补充接口调用记录、状态回读和异常处理代码。VLA部分只选择以下对象和动作，均可从当前项目已核验数据源中找到相同或近似的机器人示教。")
    midterm = [
        ("VLA-01", "杯/马克杯", "从固定台面抓取并放到指定位置", "RoboCasa：CoffeeSetupMug / CoffeeServeMug；BEHAVIOR餐具任务", "单物体、位置明确、抓取与放置可独立判定"),
        ("VLA-02", "柜门", "打开或关闭柜门", "RoboCasa：OpenCabinet / CloseCabinet", "单对象、单状态变化，不需要物体分类"),
        ("VLA-03", "微波炉", "按开始或停止按钮", "RoboCasa：TurnOnMicrowave / TurnOffMicrowave", "固定按钮、短轨迹，可采用状态谓词判定"),
        ("VLA-04", "水瓶或罐装物", "从台面放入柜体", "RoboCasa：PickPlaceCounterToCabinet", "单件取放、目标容器明确，可测放置成功"),
        ("软件-01", "灯光/空调/电视等", "通过智慧家居接口开关或设定参数", "长虹 APP 授权接口或本地仿真桩", "不涉及机械臂；以状态一致率和响应时延验收"),
    ]
    add_table(doc, ["编号", "对象", "最小动作", "现有数据/代码依据", "选择理由"], midterm, [1.5,2.0,4.0,5.4,4.1])

    add_heading(doc, "四、对表5-2、表5-3、表5-4和表5-7的具体修改建议", 1)
    add_para(doc, "原版中与家电、环境、安防和场景控制相关的表格，应保留其任务和物体框架，但在“实现方式”栏明确标注为“网络控制型”。例如灯、空调、电视、门锁、摄像头、烟雾或温湿度传感器不应被描述为 VLA 抓取对象；其数据应记录为设备标识、控制命令、时间戳、回读状态和异常码。若当前无法取得长虹 APP 的授权接口，文档应写明“采用本地设备状态机或模拟接口完成验证”，不得写成真实设备控制已完成。")
    add_para(doc, "原表5.5、5.6、5.8、5.9涉及烹饪、清洁、递送和收纳的物体，应以当前已证实的数据替换。对于“物品递送”，优先保留水杯/马克杯、水瓶或罐装物、柜体、固定台面和指定放置区等明确对象；路线在仿真中写为起点区域、目标区域和终止状态，而不写未验证的跨房间路径。对于“整理收纳”，仅保留书籍、鞋子、文件夹、杯子、收纳盒等单件归位对象，删除儿童专用、节庆、车库、户外及多对象排序任务。对于“烹饪”，中期只保留微波炉按钮、烤箱/烤面包机旋钮或单件食材入柜等低接触任务；切配、煎炒、长时加热和多食材配方转入后期。")

    add_heading(doc, "五、数据、标签、代码与验收指标的对应关系", 1)
    add_para(doc, "每项 VLA 任务必须建立可追溯记录：任务编号、对象名称、原始英文指令、数据集名称、源回合编号、RGB 视频路径、动作 Parquet 路径、动作维度、转换状态、评测脚本与成功谓词。RoboCasa365 的共享视频与动作块需要按 episode 切分为对象专属样本；切分后目录中至少包含 RGB 视频、actions.parquet、instruction.json、source_manifest.json 和 conversion_status.json。BEHAVIOR-1K、HABIT 仅在其 RGB、动作和指令三者配对时才可进入训练库。")
    mapping = [
        ("家电/环境/安防", "设备命令、状态回读、事件日志", "APP/局域网接口或本地状态机", "命令成功率、回读一致率、响应时延、异常恢复"),
        ("单物体取放", "指令、RGB、末端动作、抓取/放置状态", "OpenVLA推理服务＋RoboCasa/LIBERO评测", "10回合任务成功率、抓取率、放置率、平均推理时延"),
        ("柜门/按钮", "指令、RGB、动作、开闭或开关状态", "OpenVLA推理＋场景成功谓词", "10回合状态改变成功率、错误触发率、时延"),
    ]
    add_table(doc, ["任务轨道", "最小数据标签", "代码/接口", "验收指标"], mapping, [3.2,5.2,4.5,4.1])

    add_heading(doc, "六、后续实施流程与阶段门", 1)
    steps = [
        "阶段0：冻结中期任务清单。网络控制型任务与 VLA 机械臂任务分表管理；每个任务只保留一个清晰成功定义。",
        "阶段1：完成数据证据核验。对每个 VLA 对象核验语言指令、RGB、机器人动作和源回合；对每个网络控制任务核验接口、设备标识、状态字段和回读方式。",
        "阶段2：建立最小闭环。先以柜门开合、微波炉按钮、杯子放置和水瓶/罐装物入柜四项运行 10 回合评测；任何一项失败先修复环境或动作映射，不扩展复杂任务。",
        "阶段3：固化指标。记录随机种子、回合数、成功谓词、失败类型、模型版本和推理时延；VLA与网络控制结果分别报告。",
        "阶段4：中期验收展示。展示任务指令、相机画面、模型输出动作、仿真状态变化及结果 JSON；网络控制任务展示命令发送、状态回读和日志。",
        "阶段5：后期扩展。仅在最小任务稳定后，再增加多物体整理、食材切配、清洁轨迹和复杂场景联动，并为每项新增能力补齐示教与评测。",
    ]
    for item in steps:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.line_spacing = 1.35
        run = p.add_run(item)
        run.font.name = "Microsoft YaHei"; run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei"); run.font.size = Pt(10.5)

    add_heading(doc, "七、结论与文档修改边界", 1)
    add_para(doc, "第五章后续修改应从“覆盖尽可能多的家庭任务”调整为“每项任务均具备可证明的数据、代码和验收路径”。中期阶段可保留原版中易于通过网络接口控制的家电、环境、安防和联动任务，同时用少量单物体 VLA 闭环证明视觉到机器人动作的能力。复杂烹饪、连续清洁、多物体整理和跨区域递送不应在无充分示教、规划和评测证据时写入中期结果。待120对象完成对象级归档并完成稳定评测后，再将其扩展为后期全屋技能库。")
    add_para(doc, "本说明不改动原版或修改版文档，不新增未经实际运行获得的性能数值。后续若需要将本说明中的建议回填到第五章，应先以真实任务清单、接口日志、数据溯源文件和评测报告为依据逐项更新。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
