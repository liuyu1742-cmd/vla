from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

root=Path(r'C:\OpenVLA-Simulator\outputs\evidence_revised_report')
p=root/'基于视觉的人类动作捕获_解析和技能学习迁移技术研究_真实证据修订版.docx'
fig=root/'figure_4_4_rollout.png'
d=Document(p)
anchor=next(x for x in d.paragraphs if x.text.startswith('案例图4-4来自seed 101'))
e=OxmlElement('w:p'); anchor._p.addnext(e); pic=Paragraph(e,anchor._parent); pic.alignment=WD_ALIGN_PARAGRAPH.CENTER; pic.add_run().add_picture(str(fig),width=Inches(6.0))
e2=OxmlElement('w:p'); pic._p.addnext(e2); cap=Paragraph(e2,pic._parent); cap.style=next(x.style for x in d.paragraphs if x.text.startswith('图 4.3')); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.add_run('图 4-4 seed 101：RoboCasa整理入柜任务的实际仿真rollout关键帧')
d.save(p)
print(p)
