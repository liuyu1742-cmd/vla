"""Audit public/local evidence for the 8-task, 60-object midterm subset."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path(r"C:\RobotProject\RobotProject\docs\家庭服务任务物体数据表.xlsx")
DEFAULT_OUTPUT = ROOT / "outputs" / "midterm_testing" / "data_audit"
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}

SELECTION = {
    "cleaning": (
        "cup", "mug", "bowl", "plate", "spoon", "fork",
        "pot", "faucet", "soap", "robot_vacuum", "sink", "carpet",
    ),
    "organizing": (
        "storage_box", "bookshelf", "wardrobe", "drawer",
        "cabinet", "desk", "coffee_table", "toy",
    ),
    "smart_cooking": (
        "rice_cooker", "induction_cooker", "oven", "microwave",
        "frying_pan", "electric_kettle", "blender", "pressure_cooker",
    ),
    "appliance_management": (
        "ceiling_light", "reading_lamp", "television", "air_conditioner",
        "refrigerator", "washing_machine", "electric_curtain", "remote_control",
    ),
    "security_monitoring": (
        "smart_lock", "security_camera", "fire_alarm", "fire_extinguisher",
        "doorknob", "padlock", "doorbell_button", "video_doorbell",
    ),
    "laundry": (
        "child_clothing", "shoes", "shirt", "pants", "towel", "wool_garment",
    ),
    "food_serving": (
        "water_cup", "wine_glass", "bottle", "food_tray", "dinner_plate",
    ),
    "object_fetching": (
        "remote_control", "mobile_phone", "book", "keys", "wallet",
    ),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def source_rows(path: Path) -> dict[tuple[str, str], dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["任务物体数据表"]
    headers = [str(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows: dict[tuple[str, str], dict[str, object]] = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        key = (str(row["task_id"]), str(row["object"]))
        rows[key] = row
    workbook.close()
    return rows


def candidate_directories(raw: object) -> list[Path]:
    values = [item.strip() for item in str(raw or "").split("|") if item.strip()]
    candidates: list[Path] = []
    for value in values:
        path = Path(value)
        if path.name.lower() == "images":
            candidates.append(path)
        else:
            candidates.append(path / "images")
            candidates.append(path)
    result: list[Path] = []
    seen: set[str] = set()
    for path in candidates:
        key = str(path).lower()
        if key not in seen:
            seen.add(key)
            result.append(path)
    return result


def image_files(directory: Path) -> list[Path]:
    if not directory.is_dir():
        return []
    return sorted(
        path
        for path in directory.rglob("*")
        if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES
    )


def evenly_spaced(items: list[Path], count: int) -> list[Path]:
    if len(items) <= count:
        return items
    if count == 1:
        return [items[len(items) // 2]]
    indexes = [round(index * (len(items) - 1) / (count - 1)) for index in range(count)]
    return [items[index] for index in indexes]


def inspect_samples(paths: Iterable[Path]) -> tuple[list[dict[str, object]], list[str]]:
    samples: list[dict[str, object]] = []
    errors: list[str] = []
    for path in paths:
        try:
            with Image.open(path) as image:
                image.verify()
            with Image.open(path) as image:
                width, height = image.size
                mode = image.mode
            samples.append(
                {
                    "path": str(path.resolve()),
                    "sha256": sha256(path),
                    "width": int(width),
                    "height": int(height),
                    "mode": mode,
                }
            )
        except Exception as exc:
            errors.append(f"{path}: {type(exc).__name__}: {exc}")
    return samples, errors


def run(source: Path, output: Path, sample_count: int) -> dict[str, object]:
    if sample_count < 10:
        raise ValueError("sample_count must be at least 10")
    rows = source_rows(source)
    selected_keys = [
        (task_id, object_id)
        for task_id, object_ids in SELECTION.items()
        for object_id in object_ids
    ]
    if len(selected_keys) != 60:
        raise AssertionError(f"selection has {len(selected_keys)} objects, expected 60")

    results: list[dict[str, object]] = []
    for index, key in enumerate(selected_keys, start=1):
        if key not in rows:
            raise KeyError(f"authoritative source is missing {key[0]}::{key[1]}")
        source_row = rows[key]
        candidates = candidate_directories(source_row["object_dir"])
        chosen = None
        files: list[Path] = []
        for candidate in candidates:
            current = image_files(candidate)
            if len(current) > len(files):
                chosen = candidate
                files = current
        sample_paths = evenly_spaced(files, sample_count)
        samples, errors = inspect_samples(sample_paths)
        duplicate_hashes = len(samples) - len({item["sha256"] for item in samples})
        passed = (
            len(files) >= sample_count
            and len(samples) == sample_count
            and not errors
            and duplicate_hashes == 0
        )
        result = {
            "index": index,
            "relation_key": f"{key[0]}::{key[1]}",
            "task_id": key[0],
            "task_name": str(source_row["task_name"]),
            "object_id": key[1],
            "object_name": str(source_row["物体中文名称"]),
            "declared_images": int(source_row["images"] or 0),
            "declared_annotations": int(source_row["annotations"] or 0),
            "resolved_directory": str(chosen.resolve()) if chosen else None,
            "discovered_image_count": len(files),
            "sample_count": len(samples),
            "duplicate_sample_hashes": duplicate_hashes,
            "errors": errors,
            "status": "PASS" if passed else "FAIL",
            "samples": samples,
        }
        results.append(result)
        print(
            f"[{index:02d}/60] {result['relation_key']}: "
            f"{result['status']} images={len(files)} samples={len(samples)}",
            flush=True,
        )

    passed = sum(item["status"] == "PASS" for item in results)
    payload = {
        "schema_version": "midterm_public_data_audit_v1",
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "source_workbook": str(source.resolve()),
        "source_sha256": sha256(source),
        "task_count": len(SELECTION),
        "object_count": len(results),
        "sample_count_per_object": sample_count,
        "passed_objects": passed,
        "failed_objects": len(results) - passed,
        "all_passed": passed == len(results),
        "results": results,
    }
    output.mkdir(parents=True, exist_ok=True)
    (output / "report.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    with (output / "summary.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        fieldnames = [
            "index", "relation_key", "task_id", "task_name", "object_id",
            "object_name", "declared_images", "declared_annotations",
            "resolved_directory", "discovered_image_count", "sample_count",
            "duplicate_sample_hashes", "status",
        ]
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for item in results:
            writer.writerow({name: item.get(name) for name in fieldnames})
    print(
        json.dumps(
            {
                "task_count": payload["task_count"],
                "object_count": payload["object_count"],
                "passed_objects": passed,
                "failed_objects": len(results) - passed,
                "report": str((output / "report.json").resolve()),
            },
            ensure_ascii=False,
        ),
        flush=True,
    )
    return payload


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--sample-count", type=int, default=10)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    payload = run(args.source, args.output, args.sample_count)
    return 0 if payload["all_passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
