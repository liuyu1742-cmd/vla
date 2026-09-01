from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("outputs/midterm_test_plan/中期技能学习仿真测试流程与准备方案.docx")
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"


def set_font(run, name="Calibri", size=11, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_margins(cell)


def write_cell(cell, text, bold=False, color=None, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT_BLUE)
        write_cell(cell, text, bold=True, color=DARK_BLUE, size=font_size)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            write_cell(cell, text, size=font_size)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text), size=11)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text), size=11)


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, LIGHT_GRAY)
    set_cell_margins(cell, 120, 160, 120, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(title + "："), size=10.5, bold=True, color=DARK_BLUE)
    set_font(p.add_run(text), size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), size={1: 16, 2: 13, 3: 12}[level], bold=True, color=BLUE if level < 3 else DARK_BLUE)


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        set_font(p.add_run(bold_prefix), size=11, bold=True)
        set_font(p.add_run(text[len(bold_prefix):]), size=11)
    else:
        set_font(p.add_run(text), size=11)
    return p


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("中期技能学习仿真测试流程与准备方案"), size=8.5, color="666666")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("OpenVLA-Simulator  |  中期检查工作文件"), size=8.5, color="666666")


def build():
    doc = Document()
    configure(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run("中期技能学习仿真测试流程与准备方案"), size=22, bold=True, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    set_font(p.add_run("适用范围：视觉动作解析、技能学习与仿真执行验证"), size=11, color="555555")

    add_table(doc, ["项目", "中期执行口径"], [
        ("中期目标", "至少完成 8 类任务、60 个单一物体的仿真测试并形成可核验测试证据。"),
        ("最终全量目标", "15 类任务、137 个单一物体、167 个操作条目；不在本中期文件中宣称已全量完成。"),
        ("验证主线", "公开/已有示范数据 → OpenVLA 技能策略 → RoboCasa/MuJoCo 或 LIBERO 仿真闭环 → 结果与证据归档。"),
        ("实机展示", "完成一个低风险抓取—放置演示，用于证明相机、机械臂和控制接口已具备后续实机验证条件。"),
        ("计划工期", "10–12 个工作日；如需重新训练模型、补充资产或排查环境，预留至 15 个工作日。"),
    ], [Inches(1.45), Inches(5.05)], 10)

    add_heading(doc, "1. 目的与边界")
    add_body(doc, "本文件用于组织中期技能学习测试工作，形成可复现、可审查的仿真测试证据。测试重点是验证视觉输入、语言/任务指令、动作策略和仿真环境之间的闭环可用性，不将中期结果等同于真实机器人上的大规模虚实迁移结论。")
    add_callout(doc, "中期不纳入的事项", "不要求完成 15 类、137 个物体、167 个操作条目的全量覆盖；不要求完成实实迁移、增量学习或真实机器人统计成功率。相关内容在中期后按单独计划补充。")

    add_heading(doc, "2. 测试对象、平台与证据")
    add_table(doc, ["对象", "本阶段做法", "证据"], [
        ("示范数据", "优先复用项目已有 RoboCasa/LIBERO 示范数据与已配置模型权重；如新增数据，记录数据集名称、版本、任务和样本数量。", "数据清单、版本号、任务配置。"),
        ("策略", "使用当前 OpenVLA 推理/轻量适配链路；固定模型权重、提示词格式、动作尺度和相机参数。", "权重路径、配置文件、启动命令。"),
        ("仿真", "优先使用 RoboCasa/MuJoCo 执行家政操作；对已有 LIBERO 任务保留为视觉扰动补充验证。", "环境版本、随机种子、回放视频、逐回合 JSON/CSV。"),
        ("实机", "仅进行低风险水杯或积木抓取—放置展示；设置限速、急停和人工监护。", "设备照片、演示视频、操作记录。"),
    ], [Inches(1.0), Inches(3.65), Inches(1.85)], 9.2)

    add_heading(doc, "3. 中期覆盖矩阵")
    add_body(doc, "下表给出中期拟覆盖的 8 类任务和 60 个单一物体。实际执行前由负责人核对仿真资产是否存在；若采用等价资产，必须在结果表中记录“原对象—仿真资产—操作映射”，不得隐去替换关系。")
    coverage = [
        ("1. 家电综合管理", "阅读灯、台灯、电视遥控器、空调控制器、冰箱门、洗衣机控制面板、电动窗帘开关、空气净化器面板", "8", "按压、拨动、开/关、调节"),
        ("2. 环境调节", "温度传感器、湿度传感器、窗户把手、窗帘、智能插座、空气清新器", "6", "读取、抓取、开/关、调节"),
        ("3. 安防/设施操作", "智能门锁、门把手、安防摄像头、烟雾报警器、燃气报警器、紧急按钮", "6", "触发、按压、转动、开/关"),
        ("4. 清洁服务", "扫把、簸箕、吸尘器、抹布、喷雾瓶、海绵、拖把、垃圾桶、洗碗刷、清洁剂瓶", "10", "抓取、擦拭、喷洒、放置"),
        ("5. 烹饪辅助", "碗、餐盘、水杯、筷子、勺子、锅、砧板、调味瓶、电热水壶、电饭煲面板", "10", "抓取、递送、摆放、开/关"),
        ("6. 养护管理", "浇水壶、花盆、绿植盆、滤网、肥料瓶、宠物食盆", "6", "抓取、倾倒/浇灌、取放、更换"),
        ("7. 物品递送", "药瓶、纸巾盒、手机、遥控器、衣物、零食盒、托盘、钥匙", "8", "抓取、搬运、递送、放置"),
        ("8. 整理收纳", "鞋子、玩具积木、收纳盒、书本、衣架、洗衣篮", "6", "抓取、归位、堆叠、分类"),
    ]
    add_table(doc, ["任务类别", "中期单一物体（共 60 个）", "数量", "代表操作"], coverage, [Inches(1.15), Inches(3.75), Inches(.45), Inches(1.15)], 8.4)
    add_callout(doc, "操作条目口径", "本中期测试覆盖与上述 60 个物体相对应的代表性操作，并单列记录已覆盖的操作条目数量。167 个操作条目是最终全量目标；本阶段只报告“已覆盖/未覆盖”，不以局部样例替代全量结论。")

    add_heading(doc, "4. 前置准备与冒烟检查")
    prep = [
        ("P1", "冻结版本", "记录 Git 提交号、Python/torch/仿真环境版本、模型权重文件、运行设备和显存。", "版本记录完整，能定位同一运行环境。"),
        ("P2", "固定配置", "固定任务指令、相机分辨率/视角、动作缩放、最大决策步数与随机种子。", "配置文件受版本管理；不得在正式批量运行中临时改参数。"),
        ("P3", "三任务冒烟", "分别运行抓取放置、推/擦、开关/转动三个代表任务。", "每项至少成功 1 回合，且生成日志与视频。"),
        ("P4", "资产核对", "核对 60 个对象或等价仿真资产是否可加载、可识别、可交互。", "覆盖矩阵中无“未定义资产”；替换关系可追溯。"),
        ("P5", "输出目录检查", "为每一回合设置唯一 ID，保存配置、种子、成功标志、动作轨迹、截图和视频。", "任意结果可由 ID 回放或复跑。"),
    ]
    add_table(doc, ["编号", "步骤", "操作", "通过条件"], prep, [Inches(.45), Inches(1.0), Inches(3.45), Inches(1.6)], 9.0)

    add_heading(doc, "5. 正式测试流程")
    formal = [
        ("T1", "建立测试单元", "以“任务类别—单一物体—代表操作”为最小测试单元；为每个单元确定任务指令、成功条件和最大步数。", "测试单元编号、配置文件。"),
        ("T2", "设置初始状态", "每个单元设置至少 3 种初始状态：物体位置/朝向、相机视角或光照中至少改变一项。", "状态编号与随机种子。"),
        ("T3", "执行回合", "每种初始状态至少执行 3 次；每个物体累计至少 10 回合。执行中禁止人工接管、真值状态恢复或专家动作注入。", "逐回合日志、视频。"),
        ("T4", "自动判定", "仅以仿真环境任务成功信号或预先定义的几何/状态条件判定成功；不可依据主观截图判断。", "success、timeout、collision 等状态字段。"),
        ("T5", "失败归因", "失败归入视觉识别、抓取/接触、轨迹规划、动作执行、环境资产或配置六类之一；保留首个失败帧。", "失败码、关键帧、复跑记录。"),
        ("T6", "复核与固化", "对初次不通过或日志异常单元复跑一次；配置错误可修复后从头重测，不得与旧结果混合统计。", "复核结论、最终汇总表。"),
    ]
    add_table(doc, ["步骤", "名称", "具体操作", "必存证据"], formal, [Inches(.45), Inches(1.0), Inches(3.6), Inches(1.45)], 8.8)

    add_heading(doc, "6. 判定规则与统计方法")
    add_bullet(doc, "单回合成功：达到任务配置预先定义的成功条件，且无严重碰撞、越界、失控或人工接管。")
    add_bullet(doc, "单一物体通过：该物体对应代表操作累计完成 10 回合，成功不少于 8 回合，即成功率不低于 80%。")
    add_bullet(doc, "任务类别通过：该类别中被纳入中期清单的物体均完成测试，且至少 1 个完整闭环任务满足成功率不低于 80%。")
    add_bullet(doc, "中期目标通过：8 类任务均通过、60 个单一物体均通过，且每项均具备可追溯日志与至少一段成功视频。")
    add_bullet(doc, "报告中同时给出总体成功率、分任务成功率、分物体成功率、平均决策步数、超时次数、碰撞次数和失败类别占比。")
    add_callout(doc, "统计说明", "如因资产限制替换对象，必须同时给出“原计划对象数”和“已采用等价资产数”。不得将仿真成功率描述为真实机器人 Sim-to-Real 成功率。")

    add_heading(doc, "7. 证据与归档要求")
    evidence = [
        ("配置", "任务配置、模型权重标识、环境版本、随机种子、运行命令", "每批次 1 份"),
        ("逐回合日志", "run_id、任务/物体/操作、初始状态、步数、成功标志、失败码、耗时", "每回合 1 条"),
        ("图像与视频", "起始帧、关键动作帧、成功/失败终止帧；成功回合视频", "每单元至少 1 组"),
        ("结果汇总", "60 物体 × 任务类别 × 操作 × 成功次数/总次数 × 成功率", "每日更新，最终冻结"),
        ("异常记录", "失败复现方法、根因分类、是否修复、修复后是否重测", "发现即记录"),
        ("实机演示", "设备型号、相机位置、限速/急停措施、演示步骤与视频", "至少 1 组"),
    ]
    add_table(doc, ["材料", "必须包含内容", "最低频次"], evidence, [Inches(1.05), Inches(4.2), Inches(1.25)], 9.2)

    add_heading(doc, "8. 排期与职责")
    schedule = [
        ("第 1 天", "范围冻结与覆盖矩阵", "确认 8 类、60 物体、代表操作、资产映射和负责人。", "测试负责人"),
        ("第 2 天", "平台自检与冒烟", "完成版本记录、三任务冒烟和输出目录检查。", "算法/仿真工程师"),
        ("第 3–4 天", "数据与任务配置", "准备/核对示范数据、任务指令、成功条件和初始状态。", "算法/数据工程师"),
        ("第 5–6 天", "代表任务预验证", "8 类任务逐类跑通，修复资产、动作尺度或配置问题。", "算法/仿真工程师"),
        ("第 7–9 天", "60 物体批量运行", "执行至少 600 个正式回合（60 物体 × 10 回合），持续归档。", "仿真工程师"),
        ("第 10 天", "失败复核与统计", "复跑异常、冻结最终统计表和失败归因。", "测试负责人"),
        ("第 11 天", "实机条件展示", "完成低风险抓取—放置演示并归档。", "机械臂操作员"),
        ("第 12 天", "材料汇编", "补齐截图、视频索引、测试过程和结果说明。", "测试负责人"),
    ]
    add_table(doc, ["时间", "阶段", "完成事项", "主责"], schedule, [Inches(.85), Inches(1.45), Inches(3.5), Inches(.7)], 8.8)
    add_body(doc, "工期假设：已有 OpenVLA 权重和基础仿真环境可启动，且有一张可稳定运行的 GPU。预计批量仿真消耗约 40–80 GPU 小时；如需重新微调、安装资产或修复关键环境问题，工期延长 3–5 个工作日。", bold_prefix="工期假设：")

    add_heading(doc, "9. 实机条件展示（非统计验收）")
    add_number(doc, "检查急停、限速、工作空间边界、相机固定和夹爪开合状态；安排一名操作员与一名安全监护人。")
    add_number(doc, "选择水杯或积木作为低风险对象，完成“识别目标—接近—抓取—移动—放置—回零”演示。")
    add_number(doc, "至少录制一段完整、无剪辑的成功视频，并拍摄相机、机械臂、控制终端和安全措施。")
    add_number(doc, "在材料中明确标注“条件展示”，不与仿真统计成功率合并，也不表述为实实迁移验证。")

    add_heading(doc, "10. 开工检查清单")
    for text in [
        "已确认 8 类任务、60 个物体及代表操作，并指定每项负责人。",
        "已建立物体—仿真资产—操作映射；不存在未说明的等价替代。",
        "已固定模型、环境、相机、动作尺度、最大步数和随机种子。",
        "已通过三任务冒烟检查，日志与视频均可生成。",
        "已确定每个物体 10 回合、至少 3 种初始状态的批量运行配置。",
        "已建立结果 CSV/JSON、截图、视频和异常记录的统一目录。",
        "已约定实机展示的设备、人员、安全监护和拍摄时间。",
    ]:
        add_bullet(doc, "□ " + text)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
