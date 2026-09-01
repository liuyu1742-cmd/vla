from copy import deepcopy
from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(r"C:\OpenVLA-Simulator")
SRC=ROOT/"outputs/evidence_revised_report/evidence_report_final_checked.docx"
REF=Path(r"C:\Users\sjtu101\Desktop/2. 基于视觉的人类动作捕获、解析和技能学习迁移技术研究.docx")
OUT=ROOT/"outputs/evidence_revised_report/evidence_report_chapter5_academic.docx"
FIG1=ROOT/"outputs/evidence_revised_report/figure_5_1_academic_framework.png"
FIG2=ROOT/"outputs/whole_home_vla_coverage/preview_tasks.png"
FIG3=ROOT/"outputs/formal_skill_augmented_r2_eval/probes/seed_101_frame0.png"

def find(doc,prefix):
    m=[p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if not m: raise RuntimeError(prefix)
    return m[-1]
def after(doc,p):
    ps=doc.paragraphs
    for i,x in enumerate(ps):
        if x._p is p._p:return ps[i+1]
    raise RuntimeError('paragraph')
def fmt(src,dst):
    if src._p.pPr is not None:
        q=dst._p.get_or_add_pPr();q.clear();q.append(deepcopy(src._p.pPr))
    if src.runs and dst.runs and src.runs[0]._element.rPr is not None:
        q=dst.runs[0]._element.get_or_add_rPr();q.clear();q.append(deepcopy(src.runs[0]._element.rPr))
def put(p,text,style):
    p.clear();p.add_run(text);fmt(style,p);return p
def insert_after(p,text,style):
    x=OxmlElement('w:p');p._p.addnext(x);q=p._parent.add_paragraph();q._p.getparent().remove(q._p);q._p=x;q._element=x
    if text:q.add_run(text)
    fmt(style,q);return q
def insert_before(p,text,style):
    x=OxmlElement('w:p');p._p.addprevious(x);q=p._parent.add_paragraph();q._p.getparent().remove(q._p);q._p=x;q._element=x
    if text:q.add_run(text)
    fmt(style,q);return q
def remove_between(first,last):
    n=first._p.getnext()
    while n is not last._p:
        nxt=n.getnext();n.getparent().remove(n);n=nxt
def add_fig_after(p,path,caption,body,cap):
    q=insert_after(p,'',body);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run().add_picture(str(path),width=Cm(13.4))
    c=insert_after(q,caption,cap);c.alignment=WD_ALIGN_PARAGRAPH.CENTER;return c
def make_fig():
    font_path=r"C:\Windows\Fonts\msyh.ttc"
    try:
        font=ImageFont.truetype(font_path,30); small=ImageFont.truetype(font_path,25); title=ImageFont.truetype(font_path,34)
    except OSError:
        font=small=title=ImageFont.load_default()
    im=Image.new("RGB",(1800,990),"white"); d=ImageDraw.Draw(im)
    boxes=[(55,70,425,210,"公开真实示教证据\nHABIT / BEHAVIOR-1K / DROID"),(485,70,855,210,"15类家庭服务任务\n120种对象物体"),(915,70,1285,210,"动作原语层\n抓取、放置、开合、旋转等"),(1345,70,1745,210,"场景约束层\n可达性、姿态、安全边界"),(195,475,610,650,"技能条目表示\n任务—对象—动作—约束—证据"),(695,475,1110,650,"检索与组合模块\n指令解析、候选排序、前后置约束"),(1195,475,1610,650,"RoboCasa任务执行接口\n状态确认与动作阶段衔接")]
    for x1,y1,x2,y2,text in boxes:
        d.rounded_rectangle((x1,y1,x2,y2),radius=18,fill="#EAF2F8",outline="#1F4E79",width=4)
        lines=text.split("\n"); yy=(y1+y2)/2-(len(lines)*18)
        for line in lines:
            bb=d.textbbox((0,0),line,font=small);d.text(((x1+x2-(bb[2]-bb[0]))/2,yy),line,font=small,fill="#1F1F1F");yy+=42
    arrows=[((425,140),(485,140)),((855,140),(915,140)),((1285,140),(1345,140)),((240,210),(350,475)),((610,210),(500,475)),((1100,210),(900,475)),((1500,210),(1400,475)),((610,562),(695,562)),((1110,562),(1195,562)),((1402,650),(900,800))]
    for a,b in arrows:
        d.line((a,b),fill="#1F4E79",width=5); d.polygon([(b[0],b[1]),(b[0]-17,b[1]-9),(b[0]-17,b[1]+9)],fill="#1F4E79")
    text="技能库输出：可追溯任务目录、对象操作清单与可组合执行计划";bb=d.textbbox((0,0),text,font=title);d.text(((1800-(bb[2]-bb[0]))/2,805),text,font=title,fill="#17365D")
    im.save(FIG1)

def evidence(doc,title):
    r=[]
    for t in doc.tables[18:26]:
        for row in t.rows[1:]:
            if len(row.cells)>=4 and row.cells[1].text.startswith(title):r.append((row.cells[2].text.strip(),row.cells[3].text.strip()))
    return r
def academic_body(title,base,items):
    base=base.split('在训练实施中',1)[0].strip()[:170]
    objs='、'.join(x[0] for x in items[:3]);ops='、'.join(x[1] for x in items[:3])
    body=(base+' 从技能库构建角度，本研究将该类任务表述为由任务目标、对象集合、目标区域和动作原语共同确定的操作单元，并以对象状态变化作为任务完成的判据。对象操作清单中，本类任务已登记%d项可追溯条目，代表性对象包括%s，对应操作包括%s。对每一条目，系统同时记录对象的初始位置、目标位置、可接触部位和必要的前置条件，使任务描述能够直接关联至室内场景中的具体交互过程。执行层面遵循“目标定位—动作实施—状态确认”的闭环：先依据视觉观测确定目标与周边障碍，再按操作类型完成接近、接触和位姿调整，最后检查对象是否进入目标区域、容器是否保持可用状态以及是否满足场景安全边界。该表述避免将数据处理或性能评价内容混入任务定义，而是突出任务本身的服务边界、对象关系和可复用操作结构。')%(len(items),objs,ops)
    if len(body)<350: body+='对于存在遮挡、容器开合或多对象顺序依赖的情形，技能条目还显式保留可达性与先后关系，以支持后续的检索和组合。'
    return body[:550]
def set_fields(doc):
    u=OxmlElement('w:updateFields');u.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val','true');doc.settings.element.append(u)
def main():
    make_fig();copy2(SRC,OUT);d=Document(OUT);r=Document(REF)
    rh1,rh2,rh3=find(r,'5.1 '),find(r,'5.2 '),find(r,'5.2.1 ');rb,rc=after(r,rh3),find(r,'图 5.1 ')
    h1,h2=find(d,'5.1 '),find(d,'5.2 ');remove_between(h1,h2);fmt(rh1,h1);fmt(rh2,h2)
    p=insert_after(h1,'本章面向家庭服务机器人在室内环境中的任务表达、对象组织与技能复用需求，构建由任务类别、对象物体、动作原语、场景约束和示教证据组成的技能库。与仅以动作名称罗列能力的方式不同，本研究将每一项技能限定在明确的服务场景和对象状态变化之中：任务类别说明服务目的，对象物体描述可交互实体，动作原语刻画可复用操作，场景约束则记录可达空间、容器姿态、易碎物体和设备安全等条件。',rb)
    p=insert_after(p,'在工程实现上，技能库以15类室内家庭服务任务为上层目录，以120种对象物体及其操作条目为中层索引，并以抓取、放置、开合、旋转、按压、推拉和擦拭等基础动作构成执行层。公开真实示教证据、对象操作清单与RoboCasa任务接口被纳入同一记录体系，使任务从自然语言描述到对象交互、再到动作执行具有可追溯关系。图5.1给出了本项目技能库的实际组织结构。',rb)
    q=insert_after(p,'',rb);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run().add_picture(str(FIG1),width=Cm(14.1));c=insert_after(q,'图 5.1 项目技能库总体架构',rc);c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    intro=after(d,h2);put(intro,'本研究按照家庭服务的主要目的、对象类型和交互方式，将技能库划分为15类室内任务。分类遵循任务边界明确、对象操作可追溯和跨类能力可组合的原则：食品存放与烹饪服务分别对应保存与热源附近操作，浴室用品服务与卫生清洁分别对应物品服务与表面维护，室内递送与收纳分别对应物体转移与固定位置归位。各类任务共享基础动作原语，但通过目标对象、目标区域和场景约束形成不同的技能语义。',rb)
    for i in range(1,16):
        h=find(d,f'5.2.{i} ');title=h.text.split(' ',1)[1];b=after(d,h);fmt(rh3,h);put(b,academic_body(title,b.text,evidence(d,title)),rb)
    h53=find(d,'5.3 ');fmt(rh1,h53);put(after(d,h53),'为保证技能库中的任务描述能够落到可核验的对象交互，本节将15类任务展开为120项对象操作条目。表5-2至表5-9统一采用“任务类别：对象+操作”的命名方式，其中“控制对象”给出可交互实体，“操作类型”给出对象状态变化或末端操作方式。该组织结构既避免以相近对象重复计数，也为检索模块提供了任务、对象和动作三个层面的索引。图5.2给出了项目生成的任务与对象分布。',rb)
    add_fig_after(after(d,h53),FIG2,'图 5.2 家庭服务任务与对象分布（项目生成）',rb,rc)
    h54,h55=find(d,'5.4 '),find(d,'5.5 ');remove_between(h54,h55);fmt(rh1,h54)
    p=insert_after(h54,'技能检索与动态组合用于将自然语言需求映射为技能库中的候选条目及其执行顺序。对于用户指令q，系统首先抽取任务意图、对象集合、目标区域和约束条件，形成查询表示q=(t_q,O_q,G_q,C_q)。任一技能条目s_i表示为s_i=(t_i,O_i,A_i,G_i,C_i,E_i)，其中t_i为任务类别，O_i为对象集合，A_i为动作原语序列，G_i为目标区域，C_i为场景约束，E_i为示教证据与状态信息。该表示与本章对象操作目录保持一致，能够避免仅依靠文本相似度而忽略对象与场景边界。',rb)
    e=insert_after(p,'Score(q,s_i)=λ_t·sim_t(t_q,t_i)+λ_o·J(O_q,O_i)+λ_g·J(G_q,G_i)+λ_c·sim_c(C_q,C_i)，  Σλ=1。',rb);e.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=insert_after(e,'式中，sim_t表示任务语义相似度，J(·)为对象或区域标签集合的Jaccard匹配系数，sim_c用于比较容器姿态、易碎性、热源邻近和可达性等约束条件，λ_t、λ_o、λ_g、λ_c为归一化权重。候选技能首先按Score(q,s_i)排序，再结合证据状态和当前环境观测剔除对象不可见、目标区域不可达或前置条件未满足的条目。以“将水杯放到托盘上”为例，系统需同时检索水杯定位、抓取、稳定搬运和托盘放置等条目，而不能仅返回包含“水杯”标签的单一动作。',rb)
    p=insert_after(p,'动态组合将多步骤任务表示为有向技能序列Π=(s_{i1},s_{i2},…,s_{im})。设x_k为第k步执行前的环境状态，则组合约束满足Pre(s_{ik})⊆x_k，且x_{k+1}=T(x_k,s_{ik})；其中Pre(·)表示操作的前置条件，T(·)表示由该操作引起的对象状态转移。对收纳、递送和清洁等任务，前置条件包括目标被定位、容器已打开、放置区域未占用以及路径满足安全约束等。该机制使后续步骤能够依赖前一步产生的状态，而不是将多项动作简单并列。',rb)
    e=insert_after(p,'Π*=arg max_Π Σ_k Score(q,s_{ik})－μ·Σ_k Risk(s_{ik},x_k)，  s.t. Pre(s_{ik})⊆x_k。',rb);e.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=insert_after(e,'式中Risk(·)用于量化碰撞、跌落、液体倾倒、热源接近和设备误触等风险，μ为风险惩罚系数。图5.3所示RoboCasa取放场景用于验证对象、目标区域与动作阶段在同一技能语义下的衔接。上述检索与组合方法以本项目已登记的任务、对象和操作为基础，因而其输出为可追溯的执行计划，而非脱离对象证据的抽象文本推断。',rb)
    add_fig_after(p,FIG3,'图 5.3 RoboCasa室内取放场景示例（项目输出）',rb,rc)
    fmt(rh1,h55);put(after(d,h55),'本章围绕家庭服务机器人技能库的构建与管理，建立了由任务类别、对象物体、动作原语、场景约束和示教证据组成的统一表达。首先，依据家庭室内服务的功能边界形成15类任务体系，并以120项对象操作条目将任务语义落实到可交互实体和可复用动作；其次，通过项目技能库总体架构、任务对象分布和RoboCasa场景实例，说明了目录、对象与执行接口之间的对应关系；最后，针对自然语言需求到多步骤操作的转换，给出了包含任务、对象、区域和约束匹配的检索模型，以及满足前置条件和风险约束的动态组合表达。由此，第五章将前四章形成的视觉感知、动作理解与策略执行能力组织为可索引、可组合、可追溯的技能单元，并为后续基于真实示教的数据接入、仿真验证和实体平台扩展提供统一的任务接口。',rb)
    set_fields(d);d.save(OUT);print(OUT)
if __name__=='__main__':main()
