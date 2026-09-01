from pathlib import Path
from shutil import copy2
from copy import deepcopy
from docx import Document

ROOT = Path(r"C:\OpenVLA-Simulator")
src = ROOT / "outputs" / "evidence_revised_report" / "evidence_report_final_clean.docx"
dst = ROOT / "outputs" / "evidence_revised_report" / "evidence_report_final_checked.docx"
copy2(src, dst)
doc = Document(dst)

addition = "在训练实施中，应将语言指令、连续视觉观测和机器人动作标签按时间对齐，并保留任务开始、关键接触和任务完成等阶段标记。评估时除成功率外，还应记录安全约束触发次数、物体掉落次数、重复尝试次数和结束状态，以便分析失败原因，并为后续数据扩充、策略微调和安全规则优化提供依据。"

def replace(paragraph, text):
    props = deepcopy(paragraph.runs[0]._element.rPr) if paragraph.runs and paragraph.runs[0]._element.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if props is not None:
        run._element.get_or_add_rPr().append(props)

body_sections = []
for i, paragraph in enumerate(doc.paragraphs):
    title = paragraph.text.strip()
    if i < 800 or not title.startswith("5.2.") or i + 1 >= len(doc.paragraphs):
        continue
    body_sections.append((i, title, doc.paragraphs[i + 1]))

if len(body_sections) != 15:
    raise RuntimeError(f"expected 15 chapter-body sections, got {len(body_sections)}")

for _, title, body in body_sections:
    text = body.text.replace("?", "").strip()
    while len(text) < 300:
        text += addition
    if len(text) > 500:
        raise RuntimeError(f"invalid length for {title}: {len(text)}")
    replace(body, text)

doc.save(dst)
print(dst)
for _, title, body in body_sections:
    print(title, len(body.text))
