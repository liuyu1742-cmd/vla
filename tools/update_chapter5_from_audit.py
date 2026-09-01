"""Synchronize Chapter 5's object table with validated local samples.

The dataset manifest is the source of truth.  The updater deliberately writes
the exact English source instruction beside each Chinese operation summary so
that a claimed route can always be checked against the demonstrated episode.
"""

from __future__ import annotations

import argparse
import copy
import json
import re
from collections import defaultdict
from pathlib import Path

from docx import Document


TASK_NAMES = {
    "5.2.5": "室内卫生清洁",
    "5.2.6": "烹饪与加热辅助",
    "5.2.7": "设备维护与工具使用",
    "5.2.8": "物品递送",
    "5.2.9": "整理收纳",
    "5.2.10": "储物设施开合",
    "5.2.11": "室内安装与布置",
    "5.2.12": "餐具与容器摆放",
    "5.2.13": "衣物鞋类与玄关归位",
    "5.2.14": "工作学习区服务",
    "5.2.15": "卫浴用品与个人卫生服务",
}


def _summary_operation(instruction: str, obj: str) -> str:
    text = " ".join(instruction.split())
    if "move the book from the bed onto either nightstand" in text:
        return "卧室：将书籍从床上移至床头柜。"
    if text.lower().startswith("pick ") and " place " in text:
        return "按源指令抓取并放置。"
    if text.lower().startswith("put "):
        return "按源指令放入指定容器或位置。"
    if text.lower().startswith("open "):
        return "打开目标设施。"
    if text.lower().startswith("close "):
        return "关闭目标设施。"
    if "slide" in text.lower():
        return "按源指令推入或拉出部件。"
    if text.lower().startswith("navigate "):
        return "按源指令导航至目标设备。"
    if "wash" in text.lower():
        return "按源指令完成清洗。"
    if "organize" in text.lower() or "sort" in text.lower() or "divide" in text.lower():
        return "按源指令完成分类与归位。"
    if "attach" in text.lower() or "hang" in text.lower():
        return "按源指令完成安装或悬挂。"
    return f"按源指令操作{obj}。"


def build_rows(records: list[dict]) -> list[dict]:
    rows = []
    for record in records:
        if not record.get("triple_complete"):
            continue
        code = record["task_code"]
        if code not in TASK_NAMES:
            continue
        instruction = record["instruction"]
        rows.append(
            {
                "task_code": code,
                "task_name": TASK_NAMES[code],
                "object": record["object"],
                "operation": _summary_operation(instruction, record["object"]),
                "source_instruction": instruction,
                "status": "指令、RGB 视频、动作标签已配对",
            }
        )
    return sorted(rows, key=lambda row: (tuple(map(int, row["task_code"].split("."))), row["object"]))


def add_column(table) -> None:
    """Append one physical OOXML table column and its grid definition."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    grid_col = OxmlElement("w:gridCol")
    grid_col.set(qn("w:w"), "3000")
    table._tbl.tblGrid.append(grid_col)
    for tr in table._tbl.tr_lst:
        tr.append(copy.deepcopy(tr.tc_lst[-1]))

def _set_cell(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "宋体"
            run.font.size = None


def _find_heading_index(doc: Document, code: str) -> int:
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().startswith(code + " "):
            return index
    raise ValueError(f"heading not found: {code}")


def _replace_body_after_heading(doc: Document, code: str, content: str) -> None:
    index = _find_heading_index(doc, code)
    if index + 1 >= len(doc.paragraphs):
        return
    paragraph = doc.paragraphs[index + 1]
    style = paragraph.style
    paragraph.clear()
    paragraph.style = style
    paragraph.add_run(content)


def _set_heading(doc: Document, code: str, title: str) -> None:
    paragraph = doc.paragraphs[_find_heading_index(doc, code)]
    paragraph.clear()
    paragraph.add_run(f"{code} {title}")


def _resize_object_table(table, desired_rows: int) -> None:
    while len(table.rows) < desired_rows + 1:
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
    while len(table.rows) > desired_rows + 1:
        table._tbl.remove(table.rows[-1]._tr)


def _rewrite_object_table(table, rows: list[dict]) -> None:
    _resize_object_table(table, len(rows))
    headers = ["技能名称", "操作对象", "操作说明（源指令概括）", "源指令（与视频和动作标签同源）", "核验状态"]
    while len(table.columns) < len(headers):
        add_column(table)
    for col, header in enumerate(headers):
        _set_cell(table.cell(0, col), header)
    for row_index, row in enumerate(rows, start=1):
        values = [row["task_name"], row["object"], row["operation"], row["source_instruction"], row["status"]]
        for col, value in enumerate(values):
            _set_cell(table.cell(row_index, col), value)


def _rewrite_summary_table(table, grouped: dict[str, list[dict]]) -> None:
    heading_to_code = {name: code for code, name in TASK_NAMES.items()}
    for row in table.rows[1:]:
        code = row.cells[0].text.strip()
        if code in TASK_NAMES:
            objects = "、".join(item["object"] for item in grouped[code])
            _set_cell(row.cells[1], TASK_NAMES[code])
            _set_cell(row.cells[3], f"已核验 {len(grouped[code])} 种：{objects}")
            _set_cell(row.cells[4], "仅计入指令、RGB 视频与动作标签均已配对的样本；详见表5-2。")
    _set_cell(table.cell(0, 3), "已核验对象")


def update_document(input_doc: Path, output_doc: Path, audit_json: Path) -> dict:
    audit = json.loads(audit_json.read_text(encoding="utf-8"))
    rows = build_rows(audit["records"])
    if len(rows) != 88:
        raise ValueError(f"expected 88 validated VLA objects, got {len(rows)}")
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["task_code"]].append(row)
    doc = Document(input_doc)
    _set_heading(doc, "5.2.8", "物品递送")
    _set_heading(doc, "5.2.9", "整理收纳")
    _set_heading(doc, "5.2.12", "餐具与容器摆放")
    _rewrite_summary_table(doc.tables[0], grouped)
    _rewrite_object_table(doc.tables[1], rows)
    for code, items in grouped.items():
        objects = "、".join(item["object"] for item in items)
        extra = ""
        if code == "5.2.8":
            extra = "其中书籍样本的真实路线为“床上→床头柜”；其余样本均按右列源指令的起点与终点执行，不使用旧表中的桌面→书架表述。"
        _replace_body_after_heading(
            doc,
            code,
            f"本任务当前收录{len(items)}种已核验对象：{objects}。每个对象均已逐项核验任务指令、RGB视频和机器人动作Parquet的同源配对，具体操作和源回合指令见表5-2。{extra}",
        )
    doc.save(output_doc)
    return {"validated_vla_objects": len(rows), "by_task": {key: len(value) for key, value in grouped.items()}}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--doc", type=Path, required=True)
    parser.add_argument("--audit", type=Path, required=True)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    output = args.out or args.doc
    result = update_document(args.doc, output, args.audit)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
