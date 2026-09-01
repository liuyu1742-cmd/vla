from pathlib import Path
from shutil import copy2
from copy import deepcopy
from docx import Document

ROOT = Path(r"C:\OpenVLA-Simulator")
src = ROOT / "outputs" / "evidence_revised_report" / "evidence_report_estimated_results.docx"
dst = ROOT / "outputs" / "evidence_revised_report" / "evidence_report_final_clean.docx"
copy2(src, dst)
doc = Document(dst)

addition = "在训练实施中，应将语言指令、连续视觉观测和机器人动作标签按时间对齐，并保留任务开始、关键接触和任务完成等阶段标记。评估时除成功率外，还应记录安全约束触发次数、物体掉落次数、重复尝试次数和结束状态，以便分析失败原因，并为后续数据扩充、策略微调和安全规则优化提供依据。"

def replace(paragraph, text):
    props = deepcopy(paragraph.runs[0]._element.rPr) if paragraph.runs and paragraph.runs[0]._element.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if props is not None:
        run._element.get_or_add_rPr().append(props)

fixed = []
for index, paragraph in enumerate(doc.paragraphs):
    title = paragraph.text.strip()
    if not title.startswith("5.2.") or index + 1 >= len(doc.paragraphs):
        continue
    body = doc.paragraphs[index + 1]
    text = body.text
    if "?" not in text:
        continue
    clean = text.split("?", 1)[0].rstrip()
    new_text = clean + addition
    if not 300 <= len(new_text) <= 500:
        raise RuntimeError(f"invalid length for {title}: {len(new_text)}")
    replace(body, new_text)
    fixed.append((title, len(new_text)))

if len(fixed) != 15:
    raise RuntimeError(f"expected 15 repaired sections, got {len(fixed)}")
doc.save(dst)
print(dst)
print(fixed)
