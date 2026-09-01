from pathlib import Path
from docx import Document

DOC=Path(r"outputs/chapter5_midterm_revision/第五章（中期验收版_修订_分类与公式版）.docx")
NAMES={
'5.2.5':'室内卫生清洁','5.2.6':'烹饪与加热辅助','5.2.7':'设备维护与工具使用','5.2.8':'物品递送','5.2.9':'整理收纳','5.2.10':'储物设施开合','5.2.11':'室内安装与布置','5.2.12':'餐具与容器摆放','5.2.13':'衣物鞋类与玄关归位','5.2.14':'工作学习区服务','5.2.15':'卫浴用品与个人卫生服务'}
BODY_527='设备维护与工具使用任务是指对家庭常用维护工具、设备耗材及相关储物设施进行取放、归位和状态处理的操作。家庭中的工具通常包括电钻、钳子、螺丝刀、内六角扳手、手电筒和工具箱等，它们用于完成简单维修、设备检查和日常维护。该类任务按照操作对象可分为工具归位类、工具箱管理类和设备辅助类。工具归位类要求机器人识别不同工具的外形、抓取部位和安全方向，将工具从工作台放入指定工具箱；工具箱管理类要求机器人保持工具箱位置稳定，并在工具收纳完成后关闭箱盖；设备辅助类涉及对手电筒等便携设备的摆放与取用。机器人执行时需要根据语言指令确定目标工具和收纳位置，避免工具相互碰撞或尖锐部位朝向人员活动区域，并通过视觉反馈确认工具箱内的收纳状态。'
BODY_5211='室内安装与布置任务是指对家庭室内装饰物、拍摄设备和局部设施进行摆放、悬挂或连接的操作。室内布置直接影响居住空间的整洁性和使用便利性，常见对象包括蜡烛、海报、墙钉、数码相机和相机三脚架等。该类任务按照操作方式可分为台面摆放类、墙面悬挂类和设备连接类。台面摆放类要求机器人将蜡烛等装饰物从储物位置取出并放置在指定台面区域；墙面悬挂类要求机器人将海报对准墙钉并完成悬挂；设备连接类要求机器人将数码相机安装到三脚架并保持连接稳定。机器人执行时需要识别目标区域、连接部位和物体姿态，合理规划末端移动路径，并在操作结束后通过视觉信息确认物体位置、方向和连接状态满足布置要求。'

d=Document(DOC)
# Main headings and their immediate body paragraphs.
for code,name in NAMES.items():
    idx=next(i for i,p in enumerate(d.paragraphs) if p.text.strip().startswith(code))
    d.paragraphs[idx].clear(); d.paragraphs[idx].add_run(f'{code} {name}任务')
    if code=='5.2.7': d.paragraphs[idx+1].clear(); d.paragraphs[idx+1].add_run(BODY_527)
    if code=='5.2.11': d.paragraphs[idx+1].clear(); d.paragraphs[idx+1].add_run(BODY_5211)
# Summary table.
for row in d.tables[0].rows[1:]:
    code=row.cells[0].text.strip()
    if code in NAMES: row.cells[1].text=NAMES[code]
# Eleven categorized VLA tables are the final eleven tables.
for idx,(code,name) in enumerate(NAMES.items(),1):
    table=d.tables[-11+idx-1]
    for row in table.rows[1:]: row.cells[1].text=name
    caption=f'表5-8-{idx} {name}对象操作技能（{len(table.rows)-1}种）'
    # Locate matching caption by current table ordinal and replace it.
    old_prefix=f'表5-8-{idx} '
    for p in d.paragraphs:
        if p.text.strip().startswith(old_prefix):
            p.clear(); p.add_run(caption); break
# Introduce one consistent naming standard.
for i,p in enumerate(d.paragraphs):
    if p.text.strip().startswith('5.3'):
        d.paragraphs[i+1].clear(); d.paragraphs[i+1].add_run('本节按照RoboCasa/VLA数据中的真实任务类别，对非网络控制对象进行分类说明。表5-8-1至表5-8-11的技能名称与第5.2.5至第5.2.15节完全一致，均采用同一任务名称；各表依次列出序号、技能名称、操作对象和依据源指令归纳的操作类型。')
        break
d.save(DOC); print('names_aligned')
