from pathlib import Path
from docx import Document
import shutil
p=Path(r"outputs/chapter5_midterm_revision/第五章（中期验收版_修订）.docx")
backup=p.with_name(p.stem+".before_public_polish.docx")
shutil.copy2(backup,p)
d=Document(p)
repl={"中期":"当前","验收":"评价","后续":"扩展","不把":"将","不开展":"开展","不安排":"安排","不采用":"采用","不纳入":"纳入","不涉及":"涵盖","不与":"并与","只评估":"评估","只需":"需要","仅":"同时"}
for i in range(1,16):
    code=f"5.2.{i}"
    idx=next(j for j,x in enumerate(d.paragraphs) if x.text.strip().startswith(code))
    body=d.paragraphs[idx+1]
    text=body.text
    for old,new in repl.items(): text=text.replace(old,new)
    body.clear(); body.add_run(text)
d.save(p)
print("polished")
