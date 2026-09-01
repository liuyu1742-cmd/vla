from pathlib import Path
from shutil import copy2
from docx import Document

SRC=Path(r"outputs/chapter5_midterm_revision/第五章（中期验收版_修订_分类与公式版）.docx")
OUT=Path(r"outputs/chapter5_midterm_revision/第五章（中期验收版_修订_养护保留版）.docx")
copy2(SRC,OUT)
NAMES={
'5.2.5':'清洁用品取放与归位','5.2.6':'厨房器具单件归位','5.2.8':'物品递送','5.2.9':'整理收纳','5.2.10':'储物设施开合','5.2.11':'室内摆放','5.2.12':'餐具与容器摆放','5.2.13':'衣物鞋类与玄关归位','5.2.14':'工作学习区服务','5.2.15':'卫浴用品与个人卫生服务'}

def remove_paragraph(p):
    p._element.getparent().remove(p._element)

def main():
    d=Document(OUT)
    # Delete only the physical VLA tool-maintenance table; original Table 5-7 remains intact.
    target=next(t for t in d.tables if len(t.rows)>1 and len(t.rows[0].cells)==4 and t.rows[1].cells[1].text.strip()=='设备维护与工具使用')
    target._tbl.getparent().remove(target._tbl)
    for p in list(d.paragraphs):
        if '设备维护与工具使用对象操作技能' in p.text:
            remove_paragraph(p)
    # The first five tables are the task summary plus original network tables.
    physical=list(d.tables[5:])
    if len(physical)!=10: raise RuntimeError(f"expected 10 remaining VLA tables, found {len(physical)}")
    captions=[p for p in d.paragraphs if p.text.startswith("表5-8-") and "对象操作技能" in p.text]
    if len(captions)!=10: raise RuntimeError(f"expected 10 VLA captions, found {len(captions)}")
    number=0
    for idx,((code,name),table,caption_p) in enumerate(zip(NAMES.items(),physical,captions),1):
        for row in table.rows[1:]:
            number+=1; row.cells[0].text=str(number); row.cells[1].text=name
        caption_p.clear(); caption_p.add_run(f"表5-8-{idx} {name}对象操作技能（{len(table.rows)-1}种）")
    for i,p in enumerate(d.paragraphs):
        if p.text.strip().startswith("5.3"):
            d.paragraphs[i+1].clear(); d.paragraphs[i+1].add_run("本节按照任务类别对非网络控制对象进行分类说明。养护管理任务采用表5-7所列的10种养护对象；表5-8-1至表5-8-10列出其余10类VLA任务对象。各表均使用序号、技能名称、操作对象和操作类型四列，其中操作类型依据对应源指令、RGB视频和机器人动作标签归纳。")
            break
    d.save(OUT)
    print('saved',OUT)
if __name__=='__main__': main()


