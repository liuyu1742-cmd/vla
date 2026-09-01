"""Extract the authoritative Chapter 5 and task-object workbook for design review.

This is a read-only source-analysis helper. It writes compact JSON/text reports into
the workspace so the project specification can be audited against exact source rows.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def normalize(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def extract_docx(docx_path: Path) -> dict[str, Any]:
    document = Document(docx_path)
    paragraphs: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = normalize(paragraph.text)
        if text:
            paragraphs.append(
                {
                    "index": index,
                    "style": normalize(paragraph.style.name if paragraph.style else ""),
                    "text": text,
                }
            )

    # Prefer the actual body heading. Matching the first ``5.x`` paragraph would
    # incorrectly select the table-of-contents entry near the start of the file.
    body_headings = [
        item
        for item in paragraphs
        if item["style"].lower().startswith("heading 1")
    ]
    chapter_candidates = [
        item
        for item in body_headings
        if "技能库构建与管理" in item["text"]
        or re.search(r"第五章", item["text"])
    ]
    chapter_start = chapter_candidates[-1]["index"] if chapter_candidates else None
    chapter_end = next(
        (
            item["index"]
            for item in body_headings
            if chapter_start is not None and item["index"] > chapter_start
        ),
        None,
    )

    if chapter_start is None:
        raise RuntimeError("Could not locate Chapter 5 heading")

    chapter_paragraphs = [
        item
        for item in paragraphs
        if item["index"] >= chapter_start
        and (chapter_end is None or item["index"] < chapter_end)
    ]

    tables: list[dict[str, Any]] = []
    for table_index, table in enumerate(document.tables):
        rows = [[normalize(cell.text) for cell in row.cells] for row in table.rows]
        tables.append({"table_index": table_index, "rows": rows})

    return {
        "path": str(docx_path),
        "paragraph_count": len(document.paragraphs),
        "nonempty_paragraph_count": len(paragraphs),
        "table_count": len(document.tables),
        "chapter5_start_paragraph": chapter_start,
        "chapter6_start_paragraph": chapter_end,
        "chapter5_paragraphs": chapter_paragraphs,
        "tables": tables,
    }


def extract_workbook(xlsx_path: Path) -> dict[str, Any]:
    workbook = load_workbook(xlsx_path, data_only=False, read_only=False)
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        rows: list[dict[str, Any]] = []
        for row_index in range(1, worksheet.max_row + 1):
            values = [
                normalize(worksheet.cell(row=row_index, column=column_index).value)
                for column_index in range(1, worksheet.max_column + 1)
            ]
            if any(values):
                rows.append(
                    {
                        "row": row_index,
                        "range": f"A{row_index}:{get_column_letter(worksheet.max_column)}{row_index}",
                        "values": values,
                    }
                )
        sheets.append(
            {
                "title": worksheet.title,
                "dimensions": worksheet.calculate_dimension(),
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "merged_ranges": [str(item) for item in worksheet.merged_cells.ranges],
                "rows": rows,
            }
        )
    return {"path": str(xlsx_path), "sheets": sheets}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    docx_report = extract_docx(args.docx)
    workbook_report = extract_workbook(args.xlsx)
    (args.output_dir / "chapter5.json").write_text(
        json.dumps(docx_report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (args.output_dir / "task_object_workbook.json").write_text(
        json.dumps(workbook_report, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    chapter_text = "\n".join(
        f"[{item['index']}] {item['style']}: {item['text']}"
        for item in docx_report["chapter5_paragraphs"]
    )
    (args.output_dir / "chapter5.txt").write_text(chapter_text, encoding="utf-8")

    workbook_lines: list[str] = []
    for sheet in workbook_report["sheets"]:
        workbook_lines.append(f"## {sheet['title']} ({sheet['dimensions']})")
        for row in sheet["rows"]:
            workbook_lines.append(
                f"{row['row']}: " + " | ".join(row["values"])
            )
    (args.output_dir / "task_object_workbook.txt").write_text(
        "\n".join(workbook_lines), encoding="utf-8"
    )

    print(
        json.dumps(
            {
                "chapter5_paragraphs": len(docx_report["chapter5_paragraphs"]),
                "docx_tables": docx_report["table_count"],
                "sheets": [
                    {
                        "title": sheet["title"],
                        "dimensions": sheet["dimensions"],
                        "nonempty_rows": len(sheet["rows"]),
                    }
                    for sheet in workbook_report["sheets"]
                ],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
