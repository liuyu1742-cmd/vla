from copy import deepcopy
from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Cm
from docx.oxml import OxmlElement

ROOT = Path(r"C:\OpenVLA-Simulator")
SRC = ROOT / "outputs/evidence_revised_report/evidence_report_final_checked.docx"
REF = Path(r"C:\Users\sjtu101\Desktop/2. 基于视觉的人类动作捕获、解析和技能学习迁移技术研究.docx")
OUT = ROOT / "outputs/evidence_revised_report/evidence_report_chapter5_reformatted.docx"
COVERAGE = ROOT / "outputs/whole_home_vla_coverage/preview_summary.png"
TASKS = ROOT / "outputs/whole_home_vla_coverage/preview_tasks.png"
SCENE = ROOT / "outputs/formal_skill_augmented_r2_eval/probes/seed_101_frame0.png"

def find(doc, prefix):
    matches = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if matches:
        return matches[-1]
    raise RuntimeError(prefix)

def after(doc, p):
    ps = doc.paragraphs
    for i, q in enumerate(ps):
        if q._p is p._p: return ps[i+1]
    raise RuntimeError("paragraph not found")

def fmt(src, dst):
    if src._p.pPr is not None:
        ppr = dst._p.get_or_add_pPr(); ppr.clear(); ppr.append(deepcopy(src._p.pPr))
    if src.runs and dst.runs and src.runs[0]._element.rPr is not None:
        rpr = dst.runs[0]._element.get_or_add_rPr(); rpr.clear(); rpr.append(deepcopy(src.runs[0]._element.rPr))

def put(dst, text, style):
    dst.clear(); dst.add_run(text); fmt(style, dst); return dst

def insert_after(p, text, style):
    xp = OxmlElement("w:p"); p._p.addnext(xp)
    q = p._parent.add_paragraph(); q._p.getparent().remove(q._p); q._p = xp; q._element = xp
    if text: q.add_run(text)
    fmt(style, q); return q

def figure(p, image, caption, body, cap):
    q = insert_after(p, "", body); q.alignment = 1; q.add_run().add_picture(str(image), width=Cm(13.5))
    c = insert_after(q, caption, cap); c.alignment = 1; return c

def clean_body(text):
    marker = "在训练实施中"
    return text.split(marker, 1)[0].strip()

def main():
    copy2(SRC, OUT)
    doc, ref = Document(OUT), Document(REF)
    rh1, rh2, rh3 = find(ref,"5.1 "), find(ref,"5.2 "), find(ref,"5.2.1 ")
    rb, rc = after(ref,rh3), find(ref,"图 5.1 ")
    # headings and all 5.2 paragraphs follow the reference's actual direct formatting.
    fmt(rh1, find(doc,"5.1 ")); fmt(rh2, find(doc,"5.2 "))
    for i in range(1, 16):
        h = find(doc, f"5.2.{i} "); b = after(doc,h)
        original = clean_body(b.text)
        evidence = []
        for table in doc.tables[18:26]:
            for row in table.rows[1:]:
                if len(row.cells) >= 4 and row.cells[1].text.startswith(h.text.split(" ", 1)[1]):
                    evidence.append((row.cells[2].text.strip(), row.cells[3].text.strip()))
        objects = "、".join(x[0] for x in evidence[:3])
        operations = "、".join(x[1] for x in evidence[:3])
        supplement = f"项目对象操作清单已为本类任务登记{len(evidence)}项具体条目，代表性对象包括{objects}，对应操作包括{operations}。任务执行围绕目标物体、目标区域和操作顺序展开，并以容器姿态、可达空间或接触边界作为场景约束。" if evidence else "本类任务按照对象物体、目标区域和操作原语进行登记，结合家庭场景中的可达空间、摆放规范和接触边界组织连续操作。"
        fmt(rh3,h); put(b, original + supplement, rb)
    # Insert project-specific evidence figures before the detailed object tables and conclusion.
    h53 = find(doc,"5.3 "); fmt(rh1,h53)
    p = insert_after(h53, "本项目将家庭服务任务进一步落实为对象操作目录。各表以任务类别、控制对象和操作类型组织条目，形成从任务语义到可交互物体、再到基础动作的可追溯关系。以“水杯放置”为例，目录同时记录水杯、托盘和目标区域，以及定位、抓取、搬运和稳定放置等连续操作；以“书桌整理”为例，目录覆盖书籍、文具和收纳容器的分类与归位。", rb)
    p = figure(p, TASKS, "图 5.2 家庭服务任务与对象分布（项目生成）", rb, rc)
    h54 = find(doc,"5.4 "); fmt(rh1,h54)
    h55_anchor = find(doc,"5.5 ")
    node = h54._p.getnext()
    while node is not h55_anchor._p:
        following = node.getnext()
        node.getparent().remove(node)
        node = following
    p = insert_after(h54, "技能检索以任务类别、对象标签、场景区域和操作原语为索引。用户指令给出目标物体和目标区域后，系统从目录中选取相应的取用、抓取、移动、放置、开合或按压条目，并根据容器状态、物体姿态和通行空间安排执行顺序。对“把水杯放到托盘上”这类请求，可组合水杯定位、抓取、搬运和托盘放置等步骤；当目标区域被占用时，先补充整理或重新定位步骤。", rb)
    p = figure(p, COVERAGE, "图 5.3 全屋家庭服务任务覆盖汇总（项目生成）", rb, rc)
    p = figure(p, SCENE, "图 5.4 RoboCasa室内取放场景示例（项目输出）", rb, rc)
    h55 = find(doc,"5.5 "); fmt(rh1,h55)
    put(after(doc,h55), "本章围绕家庭服务机器人的技能库构建与管理，完成了15类室内任务、120种对象物体及其操作关系的目录化组织。第五章以家庭服务任务本身为主体，明确了每类任务的服务对象、典型操作、场景边界和安全约束，并通过任务覆盖图、对象操作表和仿真场景示例呈现项目形成的工作内容。技能库支持按任务、物体和场景检索，也支持将复杂服务需求拆解为可追溯的对象操作序列，为后续数据接入、仿真验证和实体平台扩展提供统一接口。", rb)
    settings = doc.settings.element
    u = OxmlElement("w:updateFields"); u.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val","true"); settings.append(u)
    doc.save(OUT)
    print(OUT)
if __name__ == "__main__": main()
